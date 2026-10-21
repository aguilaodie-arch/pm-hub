import io
import json
import os
import re
import shutil
import threading
import time
import uuid
import zipfile
import requests
import yaml
from urllib.parse import quote
from contextlib import asynccontextmanager
from datetime import datetime, timezone

from fastapi import (
    FastAPI,
    Request,
    Form,
    UploadFile,
    File,
    HTTPException,
    Depends,
    Query,
)
from fastapi.responses import (
    RedirectResponse,
    HTMLResponse,
    JSONResponse,
    FileResponse,
    Response,
)
from fastapi.templating import Jinja2Templates
from sqlalchemy import func
from starlette.concurrency import run_in_threadpool
from starlette.middleware.sessions import SessionMiddleware

from .config import settings, save_config, reload_config, _DEFAULT_AC_TEMPLATE, _DEFAULT_JIRA_STORY_TEMPLATE
from .database import init_db, SessionLocal
from . import models
from . import security
from . import ai
from . import extract
from . import knowledge
from . import dingtalk
from . import jira

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "templates")
templates = Jinja2Templates(directory=TEMPLATES_DIR)

# 需求六态（任意方向手动互转，反直觉约定①）与优先级
REQ_STATUSES = {"待澄清", "待设计", "设计中", "待串讲", "研发中", "已上线"}
REQ_PRIORITIES = {"P0", "P1", "P2"}
MATERIAL_EXTS = {"png", "jpg", "jpeg", "gif", "webp", "pdf", "doc", "docx",
                 "xls", "xlsx", "csv", "ppt", "pptx", "md", "txt"}
MATERIAL_MAX_SIZE = 20 * 1024 * 1024   # 单文件 ≤20MB
MATERIAL_MAX_COUNT = 20                # 单需求材料 <20 个


def _utcnow():
    return datetime.now(timezone.utc)


# ---------- helpers ----------
def get_user(request: Request):
    uid = request.session.get("user_id")
    if not uid:
        raise HTTPException(status_code=401, detail="未登录")
    db = SessionLocal()
    u = db.query(models.User).get(uid)
    db.close()
    if not u:
        raise HTTPException(status_code=401, detail="未登录")
    return u


def admin_user(request: Request):
    u = get_user(request)
    if u.role != "admin":
        raise HTTPException(403, "仅主账号可操作")
    return u


def reject_dev(user, action="执行此操作"):
    """研发角色为只读：写接口一律拒绝（仅原型问答提问与问答反馈放行，见 I4）。"""
    if user.role == "dev":
        raise HTTPException(403, "无操作权限")


def pm_or_admin(user):
    if user.role not in ("pm", "admin"):
        raise HTTPException(403, "无操作权限")


def detect_entry(abs_dir):
    """A7 入口探测：index.html（多个取最浅层）→ start.html → 全目录唯一 .html → None。"""
    htmls = []
    for root, _, files in os.walk(abs_dir):
        for f in files:
            if f.lower().endswith(".html"):
                rel = os.path.relpath(os.path.join(root, f), abs_dir).replace("\\", "/")
                htmls.append(rel)

    def shallowest(name):
        cands = [h for h in htmls if os.path.basename(h).lower() == name]
        if not cands:
            return None
        return sorted(cands, key=lambda h: (h.count("/"), h))[0]

    if len(htmls) == 1:
        unique = htmls[0]
    else:
        unique = None
    return shallowest("index.html") or shallowest("start.html") or unique


def _dir_stats(abs_dir):
    """目录内文件总数与总字节数。"""
    count, size = 0, 0
    for root, _, files in os.walk(abs_dir):
        for f in files:
            count += 1
            try:
                size += os.path.getsize(os.path.join(root, f))
            except OSError:
                pass
    return count, size


def safe_relpath(name):
    """清洗上传文件携带的相对路径（文件夹上传时形如 文件夹/子目录/文件），防目录穿越。"""
    parts = [p for p in (name or "").replace("\\", "/").split("/") if p not in ("", ".", "..")]
    return os.path.join(*parts) if parts else None


def common_top_folder(names):
    """全部文件都位于同一个顶层文件夹下时返回该文件夹名，否则返回 None。"""
    tops = set()
    for n in names or []:
        rel = safe_relpath(n)
        if not rel:
            continue
        parts = rel.replace("\\", "/").split("/")
        if len(parts) < 2:
            return None
        tops.add(parts[0])
    return tops.pop() if len(tops) == 1 else None


def chunk_text(text, max_len=800, overlap=0):
    paras = [p.strip() for p in text.split("\n") if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if len(cur) + len(p) <= max_len:
            cur += p + "\n"
        else:
            if cur:
                chunks.append(cur.strip())
            cur = p + "\n"
    if cur:
        chunks.append(cur.strip())
    if overlap > 0 and len(chunks) > 1:
        # 相邻块前向重叠：把上一块尾部 overlap 字符拼到下一块开头
        merged = [chunks[0]]
        for i in range(1, len(chunks)):
            tail = chunks[i - 1][-overlap:]
            merged.append((tail + "\n" + chunks[i]).strip())
        chunks = merged
    return chunks


_TEXT_EXTS = {".txt", ".md", ".markdown", ".log", ".csv", ".json", ".yaml", ".yml"}


def _kb_chunk(text):
    """知识库入库切块，参数读配置 knowledge.chunk_size/chunk_overlap（D20）。"""
    try:
        size = int(settings.knowledge.get("chunk_size", 1200))
        overlap = int(settings.knowledge.get("chunk_overlap", 0))
    except (TypeError, ValueError):
        size, overlap = 1200, 0
    size = max(200, min(size, 4000))
    overlap = max(0, min(overlap, size // 2))
    return chunk_text(text, size, overlap)


def _read_text_attachment(abs_path, limit=1500):
    """读取文本类附件内容（白名单扩展名），非文本附件返回 None。"""
    if os.path.splitext(abs_path)[1].lower() not in _TEXT_EXTS:
        return None
    try:
        raw = open(abs_path, "rb").read(200 * 1024)
    except OSError:
        return None
    for enc in ("utf-8", "gbk"):
        try:
            return raw.decode(enc)[:limit]
        except UnicodeDecodeError:
            continue
    return None


def _user_name(u):
    return (u.display_name or u.username) if u else "未知"


def bootstrap_admin():
    db = SessionLocal()
    try:
        uname = settings.bootstrap_admin["username"]
        if not db.query(models.User).filter_by(username=uname).first():
            h, s = security.hash_password(settings.bootstrap_admin["password"])
            u = models.User(
                username=uname,
                display_name=settings.bootstrap_admin.get("display_name", "管理员"),
                password_hash=h,
                salt=s,
                role="admin",
            )
            db.add(u)
            db.commit()
            print(f"[init] 已创建初始主账号: {uname} (密码见 config.yaml)")
    finally:
        db.close()


@asynccontextmanager
async def lifespan(app: FastAPI):
    reload_config()  # 启动即加载 config.yaml（此前只在保存配置时热更新，重启后配置会丢）
    init_db()
    bootstrap_admin()
    yield


app = FastAPI(lifespan=lifespan)
app.add_middleware(SessionMiddleware, secret_key=settings.secret_key, max_age=604800)  # 会话保持 7 天（A14）


# ---------- 页面 ----------
@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse(request, "login.html")


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    if not request.session.get("user_id"):
        return RedirectResponse("/login", status_code=303)
    db = SessionLocal()
    user = db.query(models.User).get(request.session["user_id"])
    db.close()
    return templates.TemplateResponse(request, "index.html", {"user": user})


# ---------- 认证 ----------
@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...)):
    db = SessionLocal()
    u = db.query(models.User).filter_by(username=username).first()
    db.close()
    if not u or not security.verify_password(password, u.salt, u.password_hash):
        return templates.TemplateResponse(
            request, "login.html", {"error": "登录名或密码不正确"}
        )
    if u.active is False:
        return templates.TemplateResponse(
            request, "login.html", {"error": "账号已停用，请联系主账号"}
        )
    request.session["user_id"] = u.id
    return RedirectResponse("/", status_code=303)


@app.post("/register")
def register(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    display_name: str = Form(""),
):
    if len(password) < 8:
        return templates.TemplateResponse(
            request, "login.html", {"error": "密码至少 8 位"}
        )
    db = SessionLocal()
    if db.query(models.User).filter_by(username=username).first():
        db.close()
        return templates.TemplateResponse(
            request, "login.html", {"error": "登录名已存在"}
        )
    h, s = security.hash_password(password)
    u = models.User(
        username=username,
        display_name=display_name or username,
        password_hash=h,
        salt=s,
        role="pm",
    )
    db.add(u)
    db.commit()
    db.close()
    return RedirectResponse("/login?registered=1", status_code=303)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


# ---------- 项目与版本 ----------
@app.get("/api/projects")
def list_projects(user=Depends(get_user)):
    db = SessionLocal()
    ps = db.query(models.Project).order_by(models.Project.id).all()
    data = [
        {"id": p.id, "name": p.name, "description": p.description, "owner_id": p.owner_id}
        for p in ps
    ]
    db.close()
    return data


@app.post("/api/projects")
def create_project(name: str = Form(...), description: str = Form(""), user=Depends(get_user)):
    pm_or_admin(user)
    name = name.strip()
    if not name:
        raise HTTPException(400, "项目名称不能为空")
    db = SessionLocal()
    if db.query(models.Project).filter_by(name=name).first():
        db.close()
        raise HTTPException(400, "项目已存在")
    p = models.Project(name=name, description=description, owner_id=user.id)
    db.add(p)
    db.commit()
    db.refresh(p)
    data = {"id": p.id, "name": p.name}
    db.close()
    return data


@app.delete("/api/projects/{pid}")
def delete_project(pid: int, user=Depends(get_user)):
    pm_or_admin(user)
    db = SessionLocal()
    p = db.query(models.Project).get(pid)
    if not p:
        db.close()
        raise HTTPException(404, "项目不存在")
    if db.query(models.Sprint).filter_by(project_id=pid).count():
        db.close()
        raise HTTPException(400, "请先删除该项目下的全部版本")
    db.query(models.ProjectMember).filter_by(project_id=pid).delete()
    db.delete(p)
    db.commit()
    db.close()
    return {"ok": True}


@app.get("/api/projects/{pid}/sprints")
def list_sprints(pid: int, user=Depends(get_user)):
    db = SessionLocal()
    ss = db.query(models.Sprint).filter_by(project_id=pid).order_by(models.Sprint.id).all()
    sids = [s.id for s in ss]
    pub_counts = {}
    demo_counts = {}
    if sids:  # D34：原型/demo 分开计数（版本行展示「N 原型 · M demo」）
        for s_id, kind, cnt in (
            db.query(models.Publish.sprint_id, func.coalesce(models.Publish.kind, "prototype"), func.count())
            .filter(models.Publish.sprint_id.in_(sids))
            .group_by(models.Publish.sprint_id, models.Publish.kind).all()
        ):
            if kind == "demo":
                demo_counts[s_id] = cnt
            else:
                pub_counts[s_id] = cnt
    req_counts = dict(
        db.query(models.Requirement.sprint_id, func.count())
        .filter(models.Requirement.sprint_id.in_(sids)).group_by(models.Requirement.sprint_id).all()
    ) if sids else {}
    data = [
        {
            "id": s.id,
            "name": s.name,
            "status": s.status,
            "created_at": str(s.created_at),
            "archived_at": str(s.archived_at) if s.archived_at else None,
            "publish_count": pub_counts.get(s.id, 0),
            "demo_count": demo_counts.get(s.id, 0),
            "req_count": req_counts.get(s.id, 0),
        }
        for s in ss
    ]
    db.close()
    return data


@app.post("/api/projects/{pid}/sprints")
def create_sprint(pid: int, name: str = Form(...), user=Depends(get_user)):
    pm_or_admin(user)
    name = name.strip()
    if not name:
        raise HTTPException(400, "版本名称不能为空")
    db = SessionLocal()
    proj = db.query(models.Project).get(pid)
    if not proj:
        db.close()
        raise HTTPException(404, "项目不存在")
    if db.query(models.Sprint).filter_by(project_id=pid, name=name).first():
        db.close()
        raise HTTPException(400, "该项目下版本已存在")
    s = models.Sprint(project_id=pid, name=name, status="进行中")
    db.add(s)
    db.commit()
    db.refresh(s)
    data = {"id": s.id, "name": s.name}
    db.close()
    return data


@app.delete("/api/sprints/{sid}")
def delete_sprint(sid: int, user=Depends(get_user)):
    pm_or_admin(user)
    db = SessionLocal()
    s = db.query(models.Sprint).get(sid)
    if not s:
        db.close()
        raise HTTPException(404, "版本不存在")
    if db.query(models.Publish).filter_by(sprint_id=sid).count() or \
            db.query(models.Requirement).filter_by(sprint_id=sid).count():
        db.close()
        raise HTTPException(400, "请先清空该版本下的原型与需求")
    db.delete(s)
    db.commit()
    db.close()
    return {"ok": True}


# ---------- 账号管理（仅主账号） ----------
@app.get("/api/users/options")
def user_options(user=Depends(get_user)):
    """负责 PM 下拉选项（仅 role=pm 的账号，所有登录用户可用）。"""
    db = SessionLocal()
    rows = db.query(models.User).filter_by(role="pm", active=True).order_by(models.User.id).all()
    data = [{"id": u.id, "name": _user_name(u)} for u in rows]
    db.close()
    return data


@app.get("/api/users")
def list_users(user=Depends(admin_user)):
    db = SessionLocal()
    rows = db.query(models.User).order_by(models.User.id).all()
    data = [
        {
            "id": u.id,
            "username": u.username,
            "display_name": _user_name(u),
            "role": u.role,
            "active": bool(u.active) if u.active is not None else True,
            "dingtalk_userid": u.dingtalk_userid or "",
            "created_at": str(u.created_at),
        }
        for u in rows
    ]
    db.close()
    return data


@app.post("/api/users")
def create_user(
    username: str = Form(...),
    display_name: str = Form(""),
    password: str = Form(...),
    role: str = Form("pm"),
    dingtalk_userid: str = Form(""),
    user=Depends(admin_user),
):
    username = username.strip()
    if not username:
        raise HTTPException(400, "登录名不能为空")
    if len(password) < 8:
        raise HTTPException(400, "初始密码至少 8 位")
    if role not in ("admin", "pm", "dev"):
        raise HTTPException(400, "非法角色")
    db = SessionLocal()
    if db.query(models.User).filter_by(username=username).first():
        db.close()
        raise HTTPException(400, "登录名已存在")
    if role == "admin" and db.query(models.User).filter_by(role="admin").count():
        db.close()
        raise HTTPException(400, "主账号已存在，不可再创建")
    h, s = security.hash_password(password)
    u = models.User(
        username=username,
        display_name=display_name.strip() or username,
        password_hash=h,
        salt=s,
        role=role,
        dingtalk_userid=dingtalk_userid.strip(),
    )
    db.add(u)
    db.commit()
    db.refresh(u)
    data = {"id": u.id, "username": u.username, "display_name": u.display_name, "role": u.role}
    db.close()
    return data


@app.put("/api/users/{uid}")
def update_user(
    uid: int,
    display_name: str = Form(""),
    role: str = Form(""),
    password: str = Form(""),
    dingtalk_userid: str = Form(""),  # 前端编辑时总传当前值（含清空）；FastAPI 空串→None 的强转使 Form(None) 无法区分"未传"与"清空"
    user=Depends(admin_user),
):
    db = SessionLocal()
    u = db.query(models.User).get(uid)
    if not u:
        db.close()
        raise HTTPException(404, "用户不存在")
    if role:
        if role not in ("admin", "pm", "dev"):
            db.close()
            raise HTTPException(400, "非法角色")
        if u.role == "admin" and role != "admin":
            db.close()
            raise HTTPException(400, "主账号角色不可修改")
        if role == "admin" and u.role != "admin":
            db.close()
            raise HTTPException(400, "主账号已存在，不可再创建")
        u.role = role
    if display_name.strip():
        u.display_name = display_name.strip()
    if dingtalk_userid is not None:
        u.dingtalk_userid = dingtalk_userid.strip()  # 允许清空（PUT 唯一调用方 saveUser 必带此字段）
    if password:
        if len(password) < 8:
            db.close()
            raise HTTPException(400, "密码至少 8 位")
        h, s = security.hash_password(password)
        u.password_hash, u.salt = h, s
    db.commit()
    data = {"id": u.id, "username": u.username, "display_name": u.display_name, "role": u.role}
    db.close()
    return data


@app.post("/api/users/{uid}/active")
def set_user_active(uid: int, active: str = Form(...), user=Depends(admin_user)):
    """停用/启用账号（active=true/false）。主账号不可停用。"""
    db = SessionLocal()
    u = db.query(models.User).get(uid)
    if not u:
        db.close()
        raise HTTPException(404, "用户不存在")
    if u.role == "admin":
        db.close()
        raise HTTPException(400, "主账号不可停用")
    u.active = active.strip().lower() in ("1", "true", "on", "是")
    db.commit()
    db.close()
    return {"ok": True, "active": u.active}


@app.delete("/api/users/{uid}")
def delete_user(uid: int, user=Depends(admin_user)):
    db = SessionLocal()
    u = db.query(models.User).get(uid)
    if not u:
        db.close()
        raise HTTPException(404, "用户不存在")
    if u.role == "admin":
        db.close()
        raise HTTPException(400, "主账号不可删除")
    db.delete(u)
    db.commit()
    db.close()
    return {"ok": True}


# ---------- 需求池 ----------
def _next_req_no(db):
    """RQ+yyMMdd+4位当日序列（按当天已有最大序号递增）。"""
    prefix = "RQ" + datetime.now().strftime("%y%m%d")
    rows = (
        db.query(models.Requirement.req_no)
        .filter(models.Requirement.req_no.like(prefix + "%"))
        .all()
    )
    mx = 0
    for (rn,) in rows:
        try:
            mx = max(mx, int(rn[len(prefix):]))
        except (TypeError, ValueError):
            continue
    return f"{prefix}{mx + 1:04d}"


def _req_dict(r, db):
    proj = db.query(models.Project).get(r.project_id)
    sprint = db.query(models.Sprint).get(r.sprint_id) if r.sprint_id else None
    pm = db.query(models.User).get(r.pm_id) if getattr(r, "pm_id", None) else None
    return {
        "id": r.id,
        "req_no": getattr(r, "req_no", None) or "",
        "title": r.title,
        "project_id": r.project_id,
        "project_name": proj.name if proj else "",
        "sprint_id": r.sprint_id,
        "sprint_name": sprint.name if sprint else "",
        "pm_id": getattr(r, "pm_id", None),
        "pm_name": _user_name(pm) if pm else "",
        "status": r.status,
        "priority": r.priority,
        "requester": getattr(r, "requester", None) or "",
        "description": getattr(r, "description", None) or "",
        "value": r.value or "",
        "jira_key": r.jira_key or "",
        "jira_url": (r.jira_key or "").startswith("http")
            and r.jira_key or jira.issue_url(r.jira_key or ""),  # D23：key 拼链接；兼容历史手填 URL
        "jira_sync": getattr(r, "jira_sync", None) or "",
        "jira_error": getattr(r, "jira_error", None) or "",
        "planned_online_date": getattr(r, "planned_online_date", None) or "",
        "actual_online_date": getattr(r, "actual_online_date", None) or "",
        "user_story": r.user_story or "",
        "creator_id": r.creator_id,
        "attachment_count": db.query(models.RequirementAttachment)
        .filter_by(requirement_id=r.id)
        .count(),
        "created_at": str(r.created_at),
        "updated_at": str(r.updated_at or r.created_at),
    }


@app.get("/api/requirements")
def list_requirements(
    project_id: int = 0,
    sprint_id: int = 0,
    pm_id: int = 0,
    status: str = "",
    priority: str = "",
    q: str = "",
    page: int = 0,
    page_size: int = 20,
    user=Depends(get_user),
):
    db = SessionLocal()
    query = db.query(models.Requirement)
    if project_id:
        query = query.filter_by(project_id=project_id)
    if sprint_id:
        query = query.filter_by(sprint_id=sprint_id)
    if pm_id:
        query = query.filter_by(pm_id=pm_id)
    if status:
        query = query.filter_by(status=status)
    if priority:
        query = query.filter_by(priority=priority)
    if q:
        query = query.filter(
            (models.Requirement.title.contains(q))
            | (models.Requirement.req_no.contains(q))
        )
    rs = query.order_by(models.Requirement.created_at.desc(), models.Requirement.id.desc())
    if page >= 1:  # D27：需求池分页（page 从 1 起；不传 page 保持全量返回兼容旧调用）
        page_size = max(1, min(page_size, 100))
        total = rs.count()
        rows = rs.offset((page - 1) * page_size).limit(page_size).all()
        data = [_req_dict(r, db) for r in rows]
        db.close()
        return {"items": data, "total": total, "page": page, "page_size": page_size}
    data = [_req_dict(r, db) for r in rs.all()]
    db.close()
    return data


# 导出列（D31）：对齐需求字段模型 + 系统时间列
_EXPORT_COLS = [
    ("需求编号", "req_no"), ("标题", "title"), ("项目", "project_name"),
    ("版本", "sprint_name"), ("负责 PM", "pm_name"), ("状态", "status"),
    ("优先级", "priority"), ("需求方", "requester"), ("需求价值", "value"),
    ("需求描述", "description"), ("UserStory", "user_story"),
    ("JIRA", "jira_key"), ("计划上线", "planned_online_date"),
    ("实际上线", "actual_online_date"), ("创建时间", "created_at"),
    ("更新时间", "updated_at"),
]


@app.get("/api/requirements/export")
def export_requirements(
    project_id: int = 0,
    sprint_id: int = 0,
    pm_id: int = 0,
    status: str = "",
    q: str = "",
    user=Depends(get_user),
):
    """D31：按当前筛选条件导出需求池为 xlsx（不分页，导出全部匹配结果）。"""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    db = SessionLocal()
    query = db.query(models.Requirement)
    if project_id:
        query = query.filter_by(project_id=project_id)
    if sprint_id:
        query = query.filter_by(sprint_id=sprint_id)
    if pm_id:
        query = query.filter_by(pm_id=pm_id)
    if status:
        query = query.filter_by(status=status)
    if q:
        query = query.filter(
            (models.Requirement.title.contains(q))
            | (models.Requirement.req_no.contains(q))
        )
    rs = query.order_by(models.Requirement.created_at.desc(), models.Requirement.id.desc()).all()
    rows = [_req_dict(r, db) for r in rs]
    db.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "需求池"
    ws.append([label for label, _ in _EXPORT_COLS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for d in rows:
        ws.append([d.get(key) or "" for _, key in _EXPORT_COLS])
    # 列宽：按表头与内容采样自适应（中文按 2 字符宽计，封顶 60）
    for i, (label, key) in enumerate(_EXPORT_COLS, 1):
        width = len(label) * 2 + 2
        for d in rows[:200]:
            v = str(d.get(key) or "")
            vlen = sum(2 if ord(ch) > 127 else 1 for ch in v[:58])
            width = max(width, min(60, vlen + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"需求池-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"},
    )


@app.get("/api/requirements/{rid}")
def get_requirement(rid: int, user=Depends(get_user)):
    db = SessionLocal()
    r = db.query(models.Requirement).get(rid)
    if not r:
        db.close()
        raise HTTPException(404, "需求不存在")
    data = _req_dict(r, db)
    db.close()
    return data


def _parse_req_form(db, title, project_id, sprint_id, pm_id, status, priority,
                    requester, description, value,
                    jira_url, planned_online_date, user_story):
    """校验并返回字段 dict（创建/编辑共用）。"""
    title = (title or "").strip()
    if not title or not project_id or not pm_id:
        raise HTTPException(400, "请填写完整必填项")
    if len(title) > 100:
        raise HTTPException(400, "需求标题不能超过 100 字")
    if status not in REQ_STATUSES:
        raise HTTPException(400, "非法的状态值")
    if priority not in REQ_PRIORITIES:
        raise HTTPException(400, "非法的优先级")
    if not db.query(models.Project).get(project_id):
        raise HTTPException(404, "项目不存在")
    if sprint_id and not db.query(models.Sprint).get(sprint_id):
        raise HTTPException(404, "版本不存在")
    pm = db.query(models.User).get(pm_id)
    if not pm or pm.role != "pm":
        raise HTTPException(400, "负责 PM 必须是产品经理账号")
    pod = (planned_online_date or "").strip()
    if pod:
        try:
            datetime.strptime(pod, "%Y-%m-%d")
        except ValueError:
            raise HTTPException(400, "计划上线日期格式不正确（YYYY-MM-DD）")
    return {
        "title": title,
        "project_id": project_id,
        "sprint_id": sprint_id or None,
        "pm_id": pm_id,
        "status": status,
        "priority": priority,
        "requester": (requester or "").strip(),
        "description": (description or "").strip(),
        "value": (value or "").strip(),
        "jira_key": (jira_url or "").strip(),
        "planned_online_date": pod,
        "user_story": (user_story or "").strip(),
    }


@app.post("/api/requirements")
def create_requirement(
    title: str = Form(""),
    project_id: int = Form(0),
    sprint_id: int = Form(0),
    pm_id: int = Form(0),
    status: str = Form("待澄清"),
    priority: str = Form("P1"),
    requester: str = Form(""),
    description: str = Form(""),
    value: str = Form(""),
    jira_url: str = Form(""),
    planned_online_date: str = Form(""),
    user_story: str = Form(""),
    user=Depends(get_user),
):
    pm_or_admin(user)
    db = SessionLocal()
    fields = _parse_req_form(db, title, project_id, sprint_id, pm_id, status, priority,
                             requester, description, value,
                             jira_url, planned_online_date, user_story)
    r = models.Requirement(creator_id=user.id, req_no=_next_req_no(db), **fields)
    db.add(r)
    db.commit()
    db.refresh(r)
    data = _req_dict(r, db)
    db.close()
    return data


@app.put("/api/requirements/{rid}")
def update_requirement(
    rid: int,
    title: str = Form(""),
    project_id: int = Form(0),
    sprint_id: int = Form(0),
    pm_id: int = Form(0),
    status: str = Form("待澄清"),
    priority: str = Form("P1"),
    requester: str = Form(""),
    description: str = Form(""),
    value: str = Form(""),
    jira_url: str = Form(""),
    planned_online_date: str = Form(""),
    user_story: str = Form(""),
    user=Depends(get_user),
):
    pm_or_admin(user)
    db = SessionLocal()
    r = db.query(models.Requirement).get(rid)
    if not r:
        db.close()
        raise HTTPException(404, "需求不存在")
    if r.status == "已上线":  # A4：已上线只读，仅主账号可回退
        db.close()
        raise HTTPException(403, "当前状态不可编辑")
    fields = _parse_req_form(db, title, project_id, sprint_id, pm_id, status, priority,
                             requester, description, value,
                             jira_url, planned_online_date, user_story)
    for k, v in fields.items():
        setattr(r, k, v)
    db.commit()
    data = _req_dict(r, db)
    db.close()
    return data


@app.post("/api/requirements/{rid}/rollback")
def rollback_requirement(rid: int, user=Depends(get_user)):
    """已上线 → 研发中（仅主账号，A4）。"""
    if user.role != "admin":
        raise HTTPException(403, "仅主账号可回退")
    db = SessionLocal()
    r = db.query(models.Requirement).get(rid)
    if not r:
        db.close()
        raise HTTPException(404, "需求不存在")
    if r.status != "已上线":
        db.close()
        raise HTTPException(400, "仅已上线需求可回退")
    r.status = "研发中"
    db.commit()
    data = _req_dict(r, db)
    db.close()
    return data


@app.delete("/api/requirements/{rid}")
def delete_requirement(rid: int, user=Depends(get_user)):
    pm_or_admin(user)
    db = SessionLocal()
    r = db.query(models.Requirement).get(rid)
    if not r:
        db.close()
        raise HTTPException(404, "需求不存在")
    if r.creator_id != user.id and user.role != "admin":  # A5
        db.close()
        raise HTTPException(403, "仅创建人或主账号可删除")
    if r.status == "已上线":
        db.close()
        raise HTTPException(403, "已上线需求不可删除")
    # D23：删除联动——Jira 侧 Story 流转为完成类状态并追加说明（禁止真删除 issue）；
    # Jira 不可达仅记日志，不阻塞需求删除
    jira_key = (r.jira_key or "").strip()
    if jira_key and not jira_key.startswith("http") and jira.configured():
        try:
            jira.abandon_story(
                jira_key,
                f"PM Hub 需求池侧已删除该需求（{r.req_no or ''} {r.title}），本 Story 联动废弃。",
            )
        except Exception as e:
            print(f"[warn] Jira 废弃联动失败 {jira_key}: {e}")
    # 硬删除：连带材料文件与记录、预设计对话与消息
    atts = db.query(models.RequirementAttachment).filter_by(requirement_id=rid).all()
    paths = [os.path.normpath(os.path.join(_attach_root(), a.stored_path)) for a in atts]
    for a in atts:
        db.delete(a)
    conv_ids = [
        c.id for c in db.query(models.Conversation).filter_by(requirement_id=rid).all()
    ]
    if conv_ids:
        db.query(models.Message).filter(models.Message.conversation_id.in_(conv_ids)).delete(synchronize_session=False)
        db.query(models.Conversation).filter_by(requirement_id=rid).delete()
    db.delete(r)
    db.commit()
    db.close()
    for p in paths:
        if os.path.isfile(p):
            try:
                os.remove(p)
            except OSError:
                pass
    return {"ok": True}


# ---------- Jira 联动（D23） ----------
@app.post("/api/requirements/{rid}/jira-push")
def jira_push_requirement(rid: int, user=Depends(get_user)):
    """转 Jira：LLM 按模板生成用户故事 → 在全局目标项目创建 Story → 回写关联。
    失败不阻塞需求流转：jira_sync=同步失败 + jira_error 落库，可重试（同一接口）。"""
    pm_or_admin(user)
    if not settings.ai_enabled:
        raise HTTPException(400, "请先在配置中心设置 LLM 网关")
    if not jira.configured():
        raise HTTPException(400, "请先在配置中心完成 Jira 连接配置（集成 → Jira）")
    db = SessionLocal()
    r = db.query(models.Requirement).get(rid)
    if not r:
        db.close()
        raise HTTPException(404, "需求不存在")
    if r.status == "已上线":
        db.close()
        raise HTTPException(400, "已上线需求不可转 Jira")
    if r.jira_sync == "已关联" and r.jira_key:
        db.close()
        raise HTTPException(400, f"该需求已关联 Jira {r.jira_key}，请勿重复创建")

    # 组上下文：需求字段 + 文本类材料摘录（预算 4000 字）
    attach_root = _attach_root()
    atts = db.query(models.RequirementAttachment).filter_by(requirement_id=rid).all()
    att_sections = []
    att_budget = 4000
    for a in atts:
        if att_budget <= 0:
            break
        content = _read_text_attachment(os.path.join(attach_root, a.stored_path))
        if not content:
            continue
        content = content[:att_budget]
        att_budget -= len(content)
        att_sections.append(f"材料《{a.filename}》：\n{content}")
    context = (
        f"【需求编号】{r.req_no or '-'}\n"
        f"【需求名称】{r.title}\n"
        f"【需求描述】\n{r.description or '（未填写）'}\n\n"
        f"【核心价值】\n{r.value or '（未填写）'}\n\n"
        f"【UserStory】\n{r.user_story or '（未填写）'}"
    )
    if att_sections:
        context += "\n\n【需求材料摘录】\n" + "\n\n".join(att_sections)

    template = (settings.jira.get("story_template") or "").strip() or _DEFAULT_JIRA_STORY_TEMPLATE
    try:
        story = ai.chat(
            [
                {"role": "system", "content": template},
                {"role": "user", "content": context},
            ],
            max_tokens=2000,
        ).strip()
    except Exception as e:
        db.close()
        print(f"[warn] Jira 用户故事生成失败 req={rid}: {e}")
        raise HTTPException(502, "用户故事生成失败（模型调用异常），请稍后重试")

    try:
        key = jira.create_story(r.title, story)
    except Exception as e:
        r.jira_sync = "同步失败"
        r.jira_error = str(e)
        db.commit()
        db.close()
        return {"ok": False, "jira_sync": "同步失败", "error": str(e), "story": story}
    r.jira_key = key
    r.jira_sync = "已关联"
    r.jira_error = ""
    db.commit()
    db.close()
    return {"ok": True, "jira_key": key, "jira_url": jira.issue_url(key), "story": story}


# ---------- 需求材料 ----------
def _attach_root():
    root = os.path.join(os.path.dirname(settings.publish_root), "attachments")
    os.makedirs(root, exist_ok=True)
    return root


def _att_dict(a, db):
    uploader = db.query(models.User).get(a.uploader_id) if a.uploader_id else None
    return {
        "id": a.id,
        "filename": a.filename,
        "size": a.size,
        "uploader": _user_name(uploader) if uploader else "",
        "created_at": str(a.created_at),
    }


def _get_editable_requirement(db, rid):
    r = db.query(models.Requirement).get(rid)
    if not r:
        raise HTTPException(404, "需求不存在")
    if r.status == "已上线":
        raise HTTPException(403, "当前状态不可编辑")
    return r


@app.post("/api/requirements/{rid}/attachments")
async def upload_attachment(rid: int, file: UploadFile = File(...), user=Depends(get_user)):
    pm_or_admin(user)
    db = SessionLocal()
    try:
        req = _get_editable_requirement(db, rid)
        ext = os.path.splitext(file.filename or "")[1].lstrip(".").lower()
        if ext not in MATERIAL_EXTS:
            raise HTTPException(400, "不支持该文件格式")
        count = (
            db.query(models.RequirementAttachment)
            .filter_by(requirement_id=rid)
            .count()
        )
        if count >= MATERIAL_MAX_COUNT:
            raise HTTPException(400, "单条需求材料最多 20 个")
        data_bytes = await file.read()
        if len(data_bytes) > MATERIAL_MAX_SIZE:
            raise HTTPException(400, "单个文件不能超过 20MB")
        safe_name = os.path.basename(file.filename or "file")
        rel = os.path.join(str(rid), f"{uuid.uuid4().hex[:8]}_{safe_name}")
        abs_path = os.path.join(_attach_root(), rel)
        os.makedirs(os.path.dirname(abs_path), exist_ok=True)
        with open(abs_path, "wb") as f:
            f.write(data_bytes)
        a = models.RequirementAttachment(
            requirement_id=rid,
            filename=safe_name,
            stored_path=rel,
            size=len(data_bytes),
            uploader_id=user.id,
        )
        db.add(a)
        db.commit()
        db.refresh(a)
        return _att_dict(a, db)
    finally:
        db.close()


@app.get("/api/requirements/{rid}/attachments")
def list_attachments(rid: int, user=Depends(get_user)):
    db = SessionLocal()
    rows = (
        db.query(models.RequirementAttachment)
        .filter_by(requirement_id=rid)
        .order_by(models.RequirementAttachment.id.desc())
        .all()
    )
    data = [_att_dict(a, db) for a in rows]
    db.close()
    return data


@app.get("/api/attachments/{aid}/download")
def download_attachment(aid: int, user=Depends(get_user)):
    db = SessionLocal()
    a = db.query(models.RequirementAttachment).get(aid)
    db.close()
    if not a:
        raise HTTPException(404, "文件不存在")
    base = _attach_root()
    target = os.path.normpath(os.path.join(base, a.stored_path))
    if not target.startswith(base) or not os.path.isfile(target):
        raise HTTPException(404, "文件不存在")
    return FileResponse(target, filename=a.filename)


@app.delete("/api/attachments/{aid}")
def delete_attachment(aid: int, user=Depends(get_user)):
    pm_or_admin(user)
    db = SessionLocal()
    a = db.query(models.RequirementAttachment).get(aid)
    if not a:
        db.close()
        raise HTTPException(404, "文件不存在")
    req = db.query(models.Requirement).get(a.requirement_id)
    if req and req.status == "已上线":
        db.close()
        raise HTTPException(403, "当前状态不可编辑")
    target = os.path.normpath(os.path.join(_attach_root(), a.stored_path))
    db.delete(a)
    db.commit()
    db.close()
    if os.path.isfile(target):
        try:
            os.remove(target)
        except OSError:
            pass
    return {"ok": True}


# ---------- 预设计对话（Agent-1，多对话；研发不可见，反直觉约定④） ----------
def _agent_cfg():
    return settings.design_agent


def _skills(enabled_only=True):
    skills = _agent_cfg().get("skills") or []
    if enabled_only:
        return [s for s in skills if s.get("enabled", True)]
    return skills


def _skill_catalog():
    """已安装技能目录（D15：注册制）：名称+描述+指令全文注入系统提示词，由 Agent 自主判断何时调用。"""
    skills = _skills()
    if not skills:
        return ""
    lines = []
    for i, s in enumerate(skills, 1):
        desc = (s.get("description") or "").strip()
        head = f"{i}. {s.get('name', s['key'])}" + (f"——{desc}" if desc else "")
        lines.append(f"{head}\n   指令：{(s.get('prompt') or '').strip()}")
    return (
        "【已安装技能】\n以下技能已由团队安装。当用户的意图与某个技能匹配时，主动按照该技能的指令执行；"
        "用户也可以显式点名（如「用脑暴」「帮我写 PRD」）。\n"
        "产物约定：技能产出的正式交付物（PRD、BDD 用例等文档），请将每个交付物用 ```file:文件名 的代码围栏单独包裹"
        "（例如 ```file:PRD.md 换行后接文档正文），一条消息可包含多个文件；讨论过程与说明文字放在围栏之外。\n"
        + "\n".join(lines)
    )


def reject_dev_conversation(user):
    if user.role == "dev":
        raise HTTPException(403, "无操作权限")


@app.get("/api/design-agent/skills")
def design_agent_skills(user=Depends(get_user)):
    """已安装技能列表（D15 注册制，所有登录用户可查看）。"""
    return [
        {"key": s["key"], "name": s.get("name", s["key"]), "description": s.get("description", "")}
        for s in _skills(enabled_only=False)
    ]


def _persist_skills(skills):
    save_config({"design_agent": {"skills": skills}})


def _parse_skill_markdown(raw: bytes, filename: str):
    """解析 SKILL.md：YAML frontmatter（name/description）+ 正文作为技能指令。"""
    text = raw.decode("utf-8", errors="replace")
    name, description, prompt = "", "", ""
    m = re.match(r"^---\s*\r?\n(.*?)\r?\n---\s*\r?\n?", text, re.S)
    if m:
        meta = yaml.safe_load(m.group(1)) or {}
        if isinstance(meta, dict):
            name = str(meta.get("name") or "").strip()
            description = str(meta.get("description") or "").strip()
        prompt = text[m.end():].strip()
    else:
        prompt = text.strip()
    if not name:
        name = os.path.splitext(os.path.basename(filename))[0].strip()
    return name, description, prompt


@app.post("/api/design-agent/skills/install")
async def install_skill(
    file: UploadFile = File(None),
    name: str = Form(""),
    description: str = Form(""),
    prompt: str = Form(""),
    user=Depends(admin_user),
):
    """安装技能（D15）：上传 SKILL.md / 内含 SKILL.md 的 .zip，或手工填写名称+描述+指令。"""
    s_name, s_desc, s_prompt = name.strip(), description.strip(), prompt.strip()
    if file is not None and file.filename:
        raw = await file.read()
        fname = file.filename
        if fname.lower().endswith(".zip"):
            try:
                zf = zipfile.ZipFile(io.BytesIO(raw))
            except zipfile.BadZipFile:
                raise HTTPException(400, "压缩包无法解析")
            md_names = [n for n in zf.namelist() if os.path.basename(n).lower() == "skill.md"]
            if not md_names:
                raise HTTPException(400, "压缩包中未找到 SKILL.md")
            raw = zf.read(md_names[0])
            fname = "SKILL.md"
        f_name, f_desc, f_prompt = _parse_skill_markdown(raw, fname)
        s_name = s_name or f_name
        s_desc = s_desc or f_desc
        s_prompt = s_prompt or f_prompt
    if not s_name:
        raise HTTPException(400, "请填写技能名称")
    if len(s_name) > 50:
        raise HTTPException(400, "技能名称过长（≤50 字）")
    if not s_prompt:
        raise HTTPException(400, "技能指令为空：请在文件正文中填写指令，或改用表单注册")
    skills = _skills(enabled_only=False)
    if any(s.get("name") == s_name for s in skills):
        raise HTTPException(400, f"已存在同名技能「{s_name}」，如需更新请先卸载旧版本")
    key = "skill_" + uuid.uuid4().hex[:8]
    skills.append({"key": key, "name": s_name, "description": s_desc, "prompt": s_prompt, "enabled": True})
    _persist_skills(skills)
    return {"ok": True, "skill": {"key": key, "name": s_name, "description": s_desc}}


@app.delete("/api/design-agent/skills/{key}")
def uninstall_skill(key: str, user=Depends(admin_user)):
    """卸载技能（D15）：立即从 Agent-1 上下文中移除。"""
    skills = _skills(enabled_only=False)
    remain = [s for s in skills if s.get("key") != key]
    if len(remain) == len(skills):
        raise HTTPException(404, "技能不存在或已卸载")
    _persist_skills(remain)
    return {"ok": True}


def _design_system_prompt(db, req):
    proj = db.query(models.Project).get(req.project_id)
    sprint = db.query(models.Sprint).get(req.sprint_id) if req.sprint_id else None
    pm = db.query(models.User).get(req.pm_id) if req.pm_id else None
    info = (
        f"【需求信息】\n"
        f"需求编号：{req.req_no or '-'}\n"
        f"需求名称：{req.title}\n"
        f"所属项目：{proj.name if proj else ''}\n"
        f"所属版本：{sprint.name if sprint else '未纳入版本'}\n"
        f"负责 PM：{_user_name(pm) if pm else '-'}　优先级：{req.priority}　状态：{req.status}\n"
        f"需求方：{req.requester or '-'}　JIRA：{req.jira_key or '-'}\n"
        f"需求描述：{req.description or '-'}\n"
        f"核心价值：{req.value or '-'}\n"
        f"UserStory：{req.user_story or '（未填写）'}"
    )
    # 材料摘录（A8：文本类全文注入，截断；绝不写入知识库，反直觉约定③）
    att_root = _attach_root()
    atts = (
        db.query(models.RequirementAttachment)
        .filter_by(requirement_id=req.id)
        .all()
    )
    att_sections = []
    att_budget = 8000
    for a in atts:
        if att_budget <= 0:
            break
        content = _read_text_attachment(os.path.join(att_root, a.stored_path))
        if not content:
            continue
        content = content[:att_budget]
        att_budget -= len(content)
        att_sections.append(f"材料《{a.filename}》：\n{content}")
    material_text = (
        "【需求材料摘录】\n" + "\n\n".join(att_sections) if att_sections else ""
    )
    # D15：知识库检索为默认能力，不再有开关
    kb_text = ""
    try:
        q = " ".join(x for x in [req.title, req.value or "", req.user_story or ""] if x)
        vec = ai.embed([q])[0]
        hits = knowledge.query_knowledge(req.project_id, vec, with_meta=True)
        if hits:
            labeled = [
                f"《{(m or {}).get('title') or '知识条目'}》（{(m or {}).get('req_no') or '版本索引'}）：\n{d}"
                for d, m in hits
            ]
            kb_text = "【项目知识库·相关历史逻辑摘录】\n" + "\n\n".join(labeled)
    except Exception as e:
        print(f"[warn] 预设计知识库检索失败: {e}")
    if not kb_text:
        kb_text = "【项目知识库】暂无相关历史记录（知识库来源于已归档的版本摘要）。"
    desc = (_agent_cfg().get("description") or "").strip()
    desc_block = f"【核心工作流】\n{desc}\n\n" if desc else ""
    skills_text = _skill_catalog()
    return (
        f"{desc_block}"
        f"你是「{proj.name if proj else ''}」产品团队的需求预设计助手，正在协助产品经理完成需求的前期分析与设计收敛。\n"
        f"工作方式：先基于知识库中的历史业务逻辑理解上下文，再帮助产品经理发散、讨论并收敛需求边界。\n"
        f"要求：1) 优先引用知识库中的历史逻辑，引用时说明来源；2) 知识库未覆盖的内容基于通用产品经验回答，并注明「知识库无相关记录，以下为通用建议」；"
        f"3) 使用中文，回答结构化、务实，不要寒暄。\n\n{info}\n\n{material_text}\n\n{kb_text}"
        + (f"\n\n{skills_text}" if skills_text else "")
    )


def _conv_dict(c, db):
    return {
        "id": c.id,
        "requirement_id": c.requirement_id,
        "title": c.title,
        "created_by": c.created_by,
        "message_count": db.query(models.Message).filter_by(conversation_id=c.id).count(),
        "created_at": str(c.created_at),
        "updated_at": str(c.updated_at or c.created_at),
    }


def _get_conversation(db, cid):
    c = db.query(models.Conversation).get(cid)
    if not c:
        raise HTTPException(404, "对话不存在")
    return c


@app.get("/api/requirements/{rid}/conversations")
def list_conversations(rid: int, user=Depends(get_user)):
    reject_dev_conversation(user)
    db = SessionLocal()
    cs = (
        db.query(models.Conversation)
        .filter_by(requirement_id=rid)
        .order_by(models.Conversation.updated_at.desc())
        .all()
    )
    data = [_conv_dict(c, db) for c in cs]
    db.close()
    return data


@app.post("/api/requirements/{rid}/conversations")
def create_conversation(rid: int, title: str = Form(""), user=Depends(get_user)):
    reject_dev_conversation(user)
    db = SessionLocal()
    req = _get_editable_requirement(db, rid)
    c = models.Conversation(
        requirement_id=rid,
        title=title.strip() or req.title,
        created_by=user.id,
    )
    db.add(c)
    db.commit()
    db.refresh(c)
    data = _conv_dict(c, db)
    db.close()
    return data


@app.patch("/api/conversations/{cid}")
def rename_conversation(cid: int, title: str = Form(...), user=Depends(get_user)):
    reject_dev_conversation(user)
    if not title.strip():
        raise HTTPException(400, "对话标题不能为空")
    db = SessionLocal()
    c = _get_conversation(db, cid)
    c.title = title.strip()
    db.commit()
    data = _conv_dict(c, db)
    db.close()
    return data


@app.get("/api/conversations/{cid}/messages")
def list_messages(cid: int, user=Depends(get_user)):
    reject_dev_conversation(user)
    db = SessionLocal()
    _get_conversation(db, cid)
    rows = (
        db.query(models.Message)
        .filter_by(conversation_id=cid)
        .order_by(models.Message.id)
        .all()
    )
    data = [
        {
            "id": m.id,
            "role": m.role,
            "content": m.content,
            "author_id": m.author_id,
            "created_at": str(m.created_at),
        }
        for m in rows
    ]
    db.close()
    return data


@app.post("/api/conversations/{cid}/messages")
def send_message(
    cid: int,
    message: str = Form(""),
    user=Depends(get_user),
):
    reject_dev_conversation(user)
    if not settings.design_enabled:
        raise HTTPException(400, "请先在配置中心设置 LLM 网关")
    if not message.strip():
        raise HTTPException(400, "请输入内容")
    db = SessionLocal()
    c = _get_conversation(db, cid)
    req = _get_editable_requirement(db, c.requirement_id)
    history = (
        db.query(models.Message)
        .filter_by(conversation_id=cid)
        .order_by(models.Message.id)
        .all()
    )
    system = _design_system_prompt(db, req)
    msgs = [{"role": "system", "content": system}]
    for m in history[-30:]:
        msgs.append({"role": m.role, "content": m.content})
    msgs.append({"role": "user", "content": message.strip()})

    cfg = _agent_cfg()
    max_input = int(cfg.get("max_input_length") or 0)
    if max_input > 0:
        def _total(ms):
            return sum(len(x["content"]) for x in ms)
        while len(msgs) > 2 and _total(msgs) > max_input:
            msgs.pop(1)

    temperature = float(cfg.get("temperature", 0.7))
    top_p = float(cfg.get("top_p", 0.9))
    max_tokens = int(cfg.get("max_tokens", 3000))
    max_iterations = max(1, int(cfg.get("max_iterations", 1)))

    try:
        reply = ai.chat(msgs, temperature=temperature, max_tokens=max_tokens,
                        top_p=top_p, profile="design")
        for _ in range(max_iterations - 1):
            review = msgs + [
                {"role": "assistant", "content": reply},
                {"role": "user", "content": (
                    "请检查你上面的回答：是否存在遗漏、逻辑错误或可以更深入之处？"
                    "如果没有问题，仅回复「OK」；如果有，请直接输出改进后的完整回答。"
                )},
            ]
            improved = ai.chat(review, temperature=temperature, max_tokens=max_tokens,
                               top_p=top_p, profile="design")
            if improved.strip().startswith("OK"):
                break
            reply = improved
    except Exception as e:
        db.close()
        print(f"[warn] 预设计模型调用失败: {e}")
        raise HTTPException(502, "模型调用失败，请稍后重试或联系主账号")
    display = message.strip()
    db.add(models.Message(conversation_id=cid, role="user", content=display, author_id=user.id))
    db.add(models.Message(conversation_id=cid, role="assistant", content=reply, author_id=None))
    c.updated_at = _utcnow()
    db.commit()
    db.close()
    return {"reply": reply}


@app.get("/api/conversations/{cid}/export")
def export_conversation(cid: int, user=Depends(get_user)):
    reject_dev_conversation(user)
    db = SessionLocal()
    c = _get_conversation(db, cid)
    req = db.query(models.Requirement).get(c.requirement_id)
    if not req:
        db.close()
        raise HTTPException(404, "需求不存在")
    rows = (
        db.query(models.Message)
        .filter_by(conversation_id=cid)
        .order_by(models.Message.id)
        .all()
    )
    if not rows:
        db.close()
        raise HTTPException(400, "暂无可导出内容")
    proj = db.query(models.Project).get(req.project_id)
    sprint = db.query(models.Sprint).get(req.sprint_id) if req.sprint_id else None
    pm = db.query(models.User).get(req.pm_id) if req.pm_id else None
    req_no = req.req_no or ""
    title = req.title
    db.close()
    lines = [
        f"# 需求预设计 · {title}",
        "",
        f"> 导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}　导出人：{_user_name(user)}",
        "",
        "## 需求基本信息",
        "",
        f"- 需求编号：{req_no or '-'}",
        f"- 所属项目：{proj.name if proj else '-'}",
        f"- 所属版本：{sprint.name if sprint else '未纳入版本'}",
        f"- 负责 PM：{_user_name(pm) if pm else '-'}　优先级：{req.priority}　状态：{req.status}　JIRA：{req.jira_key or '-'}",
        f"- 需求方：{req.requester or '-'}",
        f"- 核心价值：{req.value or '-'}",
        f"- UserStory：{req.user_story or '（未填写）'}",
        "",
        f"## 预设计讨论记录（{c.title}）",
        "",
    ]
    for m in rows:
        if m.role == "user":
            lines += [f"### 🧑 产品经理", "", m.content, ""]
        else:
            lines += ["### 🤖 预设计助手", "", m.content, ""]
    md = "\n".join(lines)
    filename = f"{req_no or 'RQ'}-预设计.md"
    return Response(
        content=md,
        media_type="text/markdown; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# D16：产物文件围栏约定 ```file:文件名 换行 内容 ```
_ARTIFACT_RE = re.compile(r"```file:([^\n`]+)\r?\n([\s\S]*?)```")


def _parse_artifacts(text):
    """从消息文本解析产物文件围栏，返回 [(文件名, 内容)]；文件名只取 basename 防路径注入。"""
    out = []
    for m in _ARTIFACT_RE.finditer(text or ""):
        name = os.path.basename(m.group(1).strip()) or "未命名.md"
        out.append((name, m.group(2).strip("\n")))
    return out


@app.get("/api/conversations/{cid}/artifacts.zip")
def export_artifacts(cid: int, user=Depends(get_user)):
    """打包下载当前对话的全部产物（D16：解析 AI 消息中的文件围栏，不含闲聊）。"""
    reject_dev_conversation(user)
    db = SessionLocal()
    c = _get_conversation(db, cid)
    req = db.query(models.Requirement).get(c.requirement_id)
    rows = (
        db.query(models.Message)
        .filter_by(conversation_id=cid, role="assistant")
        .order_by(models.Message.id)
        .all()
    )
    req_no = (req.req_no or "RQ") if req else "RQ"
    db.close()
    artifacts = []
    for m in rows:
        artifacts.extend(_parse_artifacts(m.content))
    if not artifacts:
        raise HTTPException(400, "本对话暂无可下载的产物")
    buf = io.BytesIO()
    used = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in artifacts:
            stem, ext = os.path.splitext(name)
            final, i = name, 1
            while final in used:  # zip 内同名自动加序号，避免相互覆盖
                i += 1
                final = f"{stem}-{i}{ext}"
            used.add(final)
            zf.writestr(final, content)
    filename = f"{req_no}-产物.zip"
    return Response(
        content=buf.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# ---------- 原型发布 ----------
def _publish_dict(p, db):
    author = _user_name(p.publisher) if p.publisher else "未知"
    entry = p.entry_path
    if not entry and p.dir_path:  # 兼容迁移后 entry_path 为 NULL 的旧数据：运行时探测兜底
        abs_dir = os.path.join(settings.publish_root, p.dir_path)
        if os.path.isdir(abs_dir):
            entry = detect_entry(abs_dir)
    link = f"/p/{p.access_key}/{entry}" if entry else f"/p/{p.access_key}/"
    return {
        "id": p.id,
        "title": p.title,
        "kind": p.kind or "prototype",
        "author": author,
        "entry_path": p.entry_path,
        "file_count": p.file_count or 0,
        "total_size": p.total_size or 0,
        "created_at": str(p.created_at),
        "updated_at": str(p.updated_at or p.created_at),
        "link": link,
    }


@app.get("/api/sprints/{sid}/publishes")
def list_publishes(sid: int, user=Depends(get_user)):
    db = SessionLocal()
    ps = db.query(models.Publish).filter_by(sprint_id=sid).order_by(models.Publish.id).all()
    data = [_publish_dict(p, db) for p in ps]
    db.close()
    return data


def _get_active_sprint(db, sid, action="上传或覆盖原型"):
    sprint = db.query(models.Sprint).get(sid)
    if not sprint:
        raise HTTPException(404, "版本不存在")
    if sprint.status == "已归档":
        raise HTTPException(403, f"该版本已归档，不可{action}")
    return sprint


# ---- 原型上传的重活（解压/写盘/文本提取），同步实现，由路由经 run_in_threadpool 调用，避免阻塞事件循环 ----

def _extract_zip(abs_dir, data):
    zip_path = os.path.join(abs_dir, "_upload.zip")
    with open(zip_path, "wb") as f:
        f.write(data)
    try:
        with zipfile.ZipFile(zip_path) as z:
            z.extractall(abs_dir)
    finally:
        if os.path.exists(zip_path):
            os.remove(zip_path)


def _write_folder(abs_dir, payloads):
    """直接上传的 HTML 文件夹：按相对路径还原目录结构。payloads=[(filename, bytes)]"""
    abs_base = os.path.abspath(abs_dir)
    for filename, data in payloads:
        rel = safe_relpath(filename)
        if not rel:
            continue
        dest = os.path.normpath(os.path.join(abs_base, rel))
        if not dest.startswith(abs_base + os.sep):
            continue
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        with open(dest, "wb") as out:
            out.write(data)


def _postprocess_publish(abs_dir):
    """入口探测 + 目录统计 + 全文提取；找不到入口返回 None"""
    entry = detect_entry(abs_dir)
    if not entry:
        return None
    file_count, total_size = _dir_stats(abs_dir)
    text = extract.extract_text(abs_dir)
    return entry, file_count, total_size, text


@app.post("/api/sprints/{sid}/publish")
async def publish(
    sid: int,
    request: Request,
    title: str = Form(""),
    kind: str = Form("prototype"),
    file: UploadFile = File(None),
    files: list[UploadFile] = File(None),
    user=Depends(get_user),
):
    pm_or_admin(user)
    kind = kind if kind in ("prototype", "demo") else "prototype"  # D34
    has_zip = file is not None and bool(file.filename)
    is_html = has_zip and file.filename.lower().endswith((".html", ".htm"))  # D34：demo 支持单 HTML 文件直传
    uploads = [f for f in (files or []) if f.filename]
    if has_zip and not (file.filename.lower().endswith(".zip") or (is_html and kind == "demo")):
        raise HTTPException(400, "请上传 ZIP 压缩包（HTML demo 还可直接上传单个 .html 文件）")
    if not has_zip and not uploads:
        raise HTTPException(400, "请上传 ZIP 压缩包或 HTML 文件夹；demo 还支持单个 .html 文件")
    db = SessionLocal()
    sprint = _get_active_sprint(db, sid, "上传或覆盖原型")
    project_id = sprint.project_id
    if title:
        final_title = title
    elif has_zip:
        final_title = file.filename
    else:
        final_title = common_top_folder([u.filename for u in uploads]) or "未命名原型"
    # 同名+同类视为覆盖：复用原记录与 access_key（已分享的链接保持有效，反直觉约定②；D34：demo 与原型互不影响）
    existing = (
        db.query(models.Publish)
        .filter_by(sprint_id=sid, title=final_title, kind=kind)
        .first()
    )
    if existing:
        existing_id = existing.id
        access_key = existing.access_key
        rel_dir = existing.dir_path
        abs_dir = os.path.join(settings.publish_root, rel_dir)
        if os.path.isdir(abs_dir):
            shutil.rmtree(abs_dir)
    else:
        existing_id = None
        access_key = uuid.uuid4().hex
        rel_dir = os.path.join(str(project_id), str(sid), access_key)
        abs_dir = os.path.join(settings.publish_root, rel_dir)
    db.close()
    os.makedirs(abs_dir, exist_ok=True)

    def _cleanup_failed(detail):
        shutil.rmtree(abs_dir, ignore_errors=True)
        if existing_id:  # 覆盖失败：旧目录已删，记录一并删除，避免悬空
            db2 = SessionLocal()
            old = db2.query(models.Publish).get(existing_id)
            if old:
                db2.delete(old)
                db2.commit()
            db2.close()
        raise HTTPException(400, detail)

    if has_zip:
        up_data = await file.read()  # 网络 IO 留在事件循环
        if is_html:
            # D34：单个 HTML 文件落盘为 index.html，直接成为入口页
            await run_in_threadpool(_write_folder, abs_dir, [("index.html", up_data)])
        else:
            try:
                await run_in_threadpool(_extract_zip, abs_dir, up_data)
            except zipfile.BadZipFile:
                _cleanup_failed("上传失败，请重试（原目录已清理）")
    else:
        # 直接上传的 HTML 文件夹：先读完全部内容（网络 IO），再统一在线程池落盘
        payloads = []
        for uf in uploads:
            payloads.append((uf.filename, await uf.read()))
        await run_in_threadpool(_write_folder, abs_dir, payloads)

    processed = await run_in_threadpool(_postprocess_publish, abs_dir)
    if not processed:
        _cleanup_failed("未找到入口页面（index.html/start.html），请检查导出内容")
    entry, file_count, total_size, text = processed

    db = SessionLocal()
    if existing_id:
        pub = db.query(models.Publish).get(existing_id)
        pub.publisher_id = user.id
        pub.extracted_text = text
        pub.entry_path = entry
        pub.file_count = file_count
        pub.total_size = total_size
        pub.updated_at = _utcnow()
        db.commit()
        data = _publish_dict(pub, db)
        db.close()
        return JSONResponse({**data, "overwritten": True})
    pub = models.Publish(
        sprint_id=sid,
        publisher_id=user.id,
        title=final_title,
        kind=kind,
        access_key=access_key,
        dir_path=rel_dir,
        extracted_text=text,
        entry_path=entry,
        file_count=file_count,
        total_size=total_size,
    )
    db.add(pub)
    db.commit()
    db.refresh(pub)
    data = _publish_dict(pub, db)
    db.close()
    return JSONResponse(data)


@app.delete("/api/publishes/{pid}")
def delete_publish(pid: int, user=Depends(get_user)):
    pm_or_admin(user)
    db = SessionLocal()
    pub = db.query(models.Publish).get(pid)
    if not pub:
        db.close()
        raise HTTPException(404, "原型不存在")
    _get_active_sprint(db, pub.sprint_id, "删除原型")
    dir_path = pub.dir_path
    db.delete(pub)
    db.commit()
    db.close()
    abs_dir = os.path.join(settings.publish_root, dir_path)
    if os.path.isdir(abs_dir):
        shutil.rmtree(abs_dir, ignore_errors=True)
    return {"ok": True}


# 研发免登录直接访问原型（D2：纯静态服务，不注入任何工具条/问答代码）
@app.get("/p/{access_key}/{filepath:path}")
def serve_proto(access_key: str, filepath: str):
    db = SessionLocal()
    pub = db.query(models.Publish).filter_by(access_key=access_key).first()
    db.close()
    if not pub:
        raise HTTPException(404, "原型不存在或已被删除")
    base = os.path.abspath(os.path.join(settings.publish_root, pub.dir_path))
    if not filepath:  # 裸链接 /p/{key}/ 兜底到入口文件
        filepath = pub.entry_path or ""
    target = os.path.normpath(os.path.join(base, filepath))
    if not target.startswith(base) or not os.path.isfile(target):
        raise HTTPException(404, "文件不存在")
    return FileResponse(target)


# ---------- 原型问答（Agent-2）+ 反馈闭环 ----------
def _agent2_system_prompt():
    return (settings.agent2.get("system_prompt") or "").strip() or (
        "你是项目的研发答疑助手。仅基于提供的资料回答，注明依据来源；资料不足时明确说明，不要编造。"
    )


@app.post("/api/publishes/{pid}/ask")
def ask_publish(pid: int, question: str = Form(...), user=Depends(get_user)):
    """针对单个原型的答疑：语料 = 该原型文本 + 该版本需求记录 + 需求文本类材料（+知识库若 agent2.use_kb）。任何登录角色可提问。"""
    if not settings.ai_enabled:
        raise HTTPException(400, "请先在配置中心设置 LLM 网关")
    if not question.strip():
        raise HTTPException(400, "请输入问题")
    db = SessionLocal()
    pub = db.query(models.Publish).get(pid)
    if not pub:
        db.close()
        raise HTTPException(404, "原型不存在")
    if (pub.kind or "prototype") == "demo":  # D34：demo 为过程体验稿，不开问答
        db.close()
        raise HTTPException(400, "HTML demo 仅供查看体验，问答请针对正式原型")
    sprint = db.query(models.Sprint).get(pub.sprint_id)
    project = db.query(models.Project).get(sprint.project_id) if sprint else None
    reqs = db.query(models.Requirement).filter_by(sprint_id=pub.sprint_id).all()
    req_ids = [r.id for r in reqs]
    atts = []
    if req_ids:
        atts = (
            db.query(models.RequirementAttachment)
            .filter(models.RequirementAttachment.requirement_id.in_(req_ids))
            .all()
        )
    attach_root = _attach_root()

    sections = []
    proto_text = (pub.extracted_text or "").strip()
    sections.append(
        f"【原型《{pub.title}》页面文字内容】\n"
        + (proto_text[:10000] if proto_text else "（该原型未提取到文字内容）")
    )
    if reqs:
        lines = []
        for r in reqs:
            pm = db.query(models.User).get(r.pm_id) if r.pm_id else None
            parts = [f"需求「{r.title}」（编号：{r.req_no or '-'}，负责 PM：{_user_name(pm) if pm else '-'}，优先级：{r.priority}，状态：{r.status}"]
            if r.jira_key:
                parts[0] += f"，JIRA：{r.jira_key}"
            parts[0] += "）"
            if r.description:
                parts.append(f"需求描述：{r.description}")
            if r.value:
                parts.append(f"核心价值：{r.value}")
            if r.user_story:
                parts.append(f"UserStory：{r.user_story}")
            lines.append("　".join(parts))
        sections.append(f"【本版本（{sprint.name if sprint else ''}）需求清单】\n" + "\n".join(lines))
    att_sections = []
    att_budget = 4000
    req_map = {r.id: r.title for r in reqs}
    for a in atts:
        if att_budget <= 0:
            break
        content = _read_text_attachment(os.path.join(attach_root, a.stored_path))
        if not content:
            continue
        content = content[:att_budget]
        att_budget -= len(content)
        att_sections.append(f"材料《{a.filename}》（属于需求「{req_map.get(a.requirement_id, '?')}」）：\n{content}")
    if att_sections:
        sections.append("【需求材料摘录】\n" + "\n\n".join(att_sections))
    if settings.agent2.get("use_kb", True) and project:
        try:
            qvec = ai.embed([question.strip()])[0]
            hits = knowledge.query_knowledge(project.id, qvec, with_meta=True)
            if hits:
                labeled = [
                    f"《{(m or {}).get('title') or '知识条目'}》（{(m or {}).get('req_no') or '版本索引'}）：\n{d}"
                    for d, m in hits
                ]
                sections.append("【知识库·相关历史逻辑摘录】\n" + "\n\n".join(labeled))
        except Exception as e:
            print(f"[warn] 原型问答知识库检索失败: {e}")
    meta = {
        "requirements": len(reqs),
        "attachments": len(att_sections),
        "proto_chars": len(proto_text[:10000]),
    }

    try:
        answer = ai.chat(
            [
                {"role": "system", "content": _agent2_system_prompt()},
                {
                    "role": "user",
                    "content": "\n\n".join(sections) + f"\n\n【研发的问题】\n{question.strip()}",
                },
            ]
        )
    except Exception as e:
        db.close()
        print(f"[warn] 原型问答模型调用失败: {e}")
        raise HTTPException(502, "模型调用失败，请稍后重试或联系主账号")
    qa = models.QaLog(
        prototype_id=pid,
        question=question.strip(),
        answer=answer,
        feedback="未反馈",
        asked_by=user.id,
    )
    db.add(qa)
    db.commit()
    db.refresh(qa)
    qa_id = qa.id
    db.close()
    return JSONResponse({"answer": answer, "qa_id": qa_id, "context": meta})


def _qa_dict(q, db):
    pub = db.query(models.Publish).get(q.prototype_id)
    asker = db.query(models.User).get(q.asked_by) if q.asked_by else None
    handler = db.query(models.User).get(q.handled_by) if q.handled_by else None
    return {
        "id": q.id,
        "prototype_id": q.prototype_id,
        "prototype_title": pub.title if pub else "（原型已删除）",
        "prototype_link": f"/p/{pub.access_key}/{pub.entry_path or ''}" if pub else "",
        "question": q.question,
        "answer": q.answer,
        "feedback": q.feedback,
        "summary": q.summary or "",
        "push_status": q.push_status or "",
        "push_error": q.push_error or "",
        "asked_by": q.asked_by,
        "asked_by_name": _user_name(asker) if asker else "",
        "handled_by_name": _user_name(handler) if handler else "",
        "handled_at": str(q.handled_at) if q.handled_at else None,
        "created_at": str(q.created_at),
    }


@app.get("/api/publishes/{pid}/qa")
def list_publish_qa(pid: int, user=Depends(get_user)):
    """当前用户自己对该原型的历史问答（问答侧栏用）。"""
    db = SessionLocal()
    rows = (
        db.query(models.QaLog)
        .filter_by(prototype_id=pid, asked_by=user.id)
        .order_by(models.QaLog.id)
        .all()
    )
    data = [_qa_dict(q, db) for q in rows]
    db.close()
    return data


@app.get("/api/qa/{qid}")
def get_qa(qid: int, user=Depends(get_user)):
    """单条问答详情（升级落地页用）。权限：提问者本人或 PM/主账号。"""
    db = SessionLocal()
    q = db.query(models.QaLog).get(qid)
    if not q:
        db.close()
        raise HTTPException(404, "问答记录不存在")
    if q.asked_by != user.id and user.role not in ("admin", "pm"):
        db.close()
        raise HTTPException(403, "无查看权限")
    pub = db.query(models.Publish).get(q.prototype_id)
    data = _qa_dict(q, db)
    data["publish_id"] = q.prototype_id
    data["publish_title"] = pub.title if pub else "（原型已删除）"
    data["can_handle"] = (
        user.role in ("admin", "pm") and q.feedback == "未解决" and not q.handled_by
    )
    db.close()
    return data


@app.post("/api/qa/{qid}/escalate")
def escalate_qa(qid: int, user=Depends(get_user)):
    """D22 升级：研发对 AI 回答点「未解决」→ LLM 按 AC 模板生成摘要 → 钉钉工作通知
    推给原型上传人。推送失败不阻断升级本身（摘要与状态照常落库）。"""
    if not settings.ai_enabled:
        raise HTTPException(400, "请先在配置中心设置 LLM 网关")
    db = SessionLocal()
    q = db.query(models.QaLog).get(qid)
    if not q:
        db.close()
        raise HTTPException(404, "问答记录不存在")
    if q.feedback == "未解决":
        db.close()
        raise HTTPException(400, "该问题已转交过产品经理")
    pub = db.query(models.Publish).get(q.prototype_id)
    if not pub:
        db.close()
        raise HTTPException(404, "原型不存在")
    sprint = db.query(models.Sprint).get(pub.sprint_id)
    project = db.query(models.Project).get(sprint.project_id) if sprint else None
    asker = db.query(models.User).get(q.asked_by) if q.asked_by else None
    owner = db.query(models.User).get(pub.publisher_id) if pub.publisher_id else None
    reqs = db.query(models.Requirement).filter_by(sprint_id=pub.sprint_id).all()
    req_titles = "、".join(f"「{r.title}」" for r in reqs) or "（本版本无关联需求）"

    context = (
        f"【所属原型】《{pub.title}》\n"
        f"【所属项目/版本】{(project.name if project else '-')} / {(sprint.name if sprint else '-')}\n"
        f"【提问人】{_user_name(asker) if asker else '未知'}（角色：{asker.role if asker else '-'}）\n"
        f"【本版本关联需求】{req_titles}\n\n"
        f"【研发的问题】\n{q.question}\n\n"
        f"【AI 的回答】\n{q.answer}"
    )
    ac_template = (settings.escalate.get("ac_template") or "").strip() or _DEFAULT_AC_TEMPLATE
    try:
        summary = ai.chat(
            [
                {"role": "system", "content": ac_template},
                {"role": "user", "content": context},
            ],
            max_tokens=1500,
        )
    except Exception as e:
        db.close()
        print(f"[warn] 升级摘要生成失败: {e}")
        raise HTTPException(502, "摘要生成失败（模型调用异常），请稍后重试")
    summary = summary.strip()
    q.feedback = "未解决"
    q.summary = summary

    hub_base = (settings.dingtalk.get("hub_base_url") or "").strip().rstrip("/")
    link = (
        f"{hub_base}/#/pc/proto-view?id={pub.id}&q={q.id}"
        if hub_base
        else "（未配置 hub_base_url，请登录 PM Hub 查看）"
    )
    md_text = (
        f"### 【PM Hub】原型问答待处理\n\n"
        f"**原型**：{pub.title}（{(project.name if project else '-')} / {(sprint.name if sprint else '-')}）\n\n"
        f"**提问人**：{_user_name(asker) if asker else '未知'}\n\n"
        f"**AC 式摘要**：\n\n{summary}\n\n"
        f"[点击查看完整上下文并处理]({link})"
    )
    push_status, push_error = "已推送", ""
    try:
        if not owner:
            raise RuntimeError("原型上传人账号已不存在")
        dingtalk.send_work_notification(
            owner.dingtalk_userid, f"【PM Hub】原型问答待处理：{pub.title}", md_text
        )
    except Exception as e:
        push_status, push_error = "推送失败", str(e)
        print(f"[warn] 钉钉推送失败 qa={q.id}: {e}")
    q.push_status = push_status
    q.push_error = push_error
    db.commit()
    owner_name = _user_name(owner) if owner else ""
    db.close()
    return {
        "ok": True,
        "summary": summary,
        "push_status": push_status,
        "push_error": push_error,
        "owner_name": owner_name,
    }


@app.post("/api/qa/{qid}/handle")
def handle_qa(qid: int, user=Depends(get_user)):
    """PM/主账号把「未解决」的升级问答标记为已处理（D22 闭环）。"""
    pm_or_admin(user)
    db = SessionLocal()
    q = db.query(models.QaLog).get(qid)
    if not q:
        db.close()
        raise HTTPException(404, "问答记录不存在")
    if q.feedback != "未解决" or q.handled_by:
        db.close()
        raise HTTPException(400, "该问答不是待处理状态")
    q.handled_by = user.id
    q.handled_at = _utcnow()
    db.commit()
    db.close()
    return {"ok": True}


# ---------- 归档 ----------
@app.get("/api/sprints/archive-board")
def archive_board(user=Depends(get_user)):
    """归档报告页数据：进行中/已归档两个列表。"""
    pm_or_admin(user)
    db = SessionLocal()
    out = {"active": [], "archived": []}
    sprints = db.query(models.Sprint).order_by(models.Sprint.id).all()
    for s in sprints:
        proj = db.query(models.Project).get(s.project_id)
        archiver = db.query(models.User).get(s.archived_by) if s.archived_by else None
        item = {
            "id": s.id,
            "project_name": proj.name if proj else "",
            "name": s.name,
            "status": s.status,
            "requirement_count": db.query(models.Requirement).filter_by(sprint_id=s.id).count(),
            "publish_count": db.query(models.Publish)
            .filter(models.Publish.sprint_id == s.id, models.Publish.kind != "demo").count(),
            "created_at": str(s.created_at),
            "archived_at": str(s.archived_at) if s.archived_at else None,
            "archived_by": _user_name(archiver) if archiver else "",
            "has_summary": bool(s.summary),
        }
        out["archived" if s.status == "已归档" else "active"].append(item)
    db.close()
    return out


@app.get("/api/sprints/{sid}/archive-preview")
def archive_preview(sid: int, user=Depends(get_user)):
    pm_or_admin(user)
    db = SessionLocal()
    sprint = db.query(models.Sprint).get(sid)
    if not sprint:
        db.close()
        raise HTTPException(404, "版本不存在")
    data = {
        "id": sprint.id,
        "name": sprint.name,
        "status": sprint.status,
        "requirement_count": db.query(models.Requirement).filter_by(sprint_id=sid).count(),
        "publish_count": db.query(models.Publish)
        .filter(models.Publish.sprint_id == sid, models.Publish.kind != "demo").count(),
    }
    db.close()
    return data


# ---------- 版本归档（D19：后台任务，需求级知识条目 + 版本级索引） ----------
ARCHIVE_JOBS = {}  # sid -> {state, step, total, done, current, error}，worker 原地更新

_ENTRY_PROMPT = (
    "你是团队的知识库整理助手。请基于以下资料，为该需求生成一张结构化的「知识卡」，"
    "供团队成员日后检索与理解业务逻辑使用。\n"
    "要求：\n"
    "1) 使用以下固定小节组织内容：\n"
    "## 背景与目标\n## 方案要点\n## 业务规则与口径\n## 异常流与边界\n## 涉及页面与联动影响\n"
    "2) 只写资料中确实有的内容，某小节无资料时写「（资料未涉及）」，不要编造；\n"
    "3) 数字、口径、规则描述以资料原文为准，不要改写；\n"
    "4) 全文 800 字以内，使用中文。"
)

_INDEX_PROMPT = (
    "你是团队的知识库整理助手。请基于以下需求知识条目清单，为该版本生成「版本索引」。\n"
    "要求：\n"
    "1) 开头用 3-5 句话概述本版本的整体目标与主要内容；\n"
    "2) 之后逐行列出每个需求的一句话定位，格式为「- 编号 名称：一句话说明」；\n"
    "3) 只基于清单内容，不要编造；使用中文。"
)


def _cosine_sim(a, b):
    dot = sum(x * y for x, y in zip(a, b))
    na = sum(x * x for x in a) ** 0.5
    nb = sum(x * x for x in b) ** 0.5
    if not na or not nb:
        return 0.0
    return dot / (na * nb)


def _entry_sources(db, req, att_root, proto_retrieve):
    """组装单条知识条目的四类来源：需求字段 + 文本材料 + 预设计对话 + 关联原型片段。
    返回 (来源文本, 关联原型 id 列表)——原型 id 供 D21 问答佐证链接使用。"""
    parts = [
        f"【需求字段】\n编号：{req.req_no or '-'}　名称：{req.title}　优先级：{req.priority}\n"
        f"需求描述：{req.description or '-'}\n核心价值：{req.value or '-'}\nUserStory：{req.user_story or '-'}"
    ]
    atts = db.query(models.RequirementAttachment).filter_by(requirement_id=req.id).all()
    att_sections = []
    budget = 6000
    for a in atts:
        if budget <= 0:
            break
        content = _read_text_attachment(os.path.join(att_root, a.stored_path), limit=3000)
        if not content:
            continue
        content = content[:budget]
        budget -= len(content)
        att_sections.append(f"材料《{a.filename}》：\n{content}")
    if att_sections:
        parts.append("【需求材料摘录】\n" + "\n\n".join(att_sections))
    convs = db.query(models.Conversation).filter_by(requirement_id=req.id).all()
    conv_lines = []
    for c in convs:
        msgs = (
            db.query(models.Message)
            .filter_by(conversation_id=c.id)
            .order_by(models.Message.id)
            .all()
        )
        for m in msgs:
            who = "产品经理" if m.role == "user" else "预设计助手"
            conv_lines.append(f"{who}：{m.content}")
    if conv_lines:
        conv_text = "\n".join(conv_lines)
        if len(conv_text) > 5000:
            conv_text = "……（前文略）\n" + conv_text[-5000:]
        parts.append("【预设计对话记录】\n" + conv_text)
    proto_hits = proto_retrieve(f"{req.title} {req.value or ''} {req.user_story or ''}")
    pub_ids = []
    if proto_hits:
        parts.append("【关联原型片段】\n" + "\n\n".join(c for _, c in proto_hits))
        for pid, _ in proto_hits:
            if pid not in pub_ids:
                pub_ids.append(pid)
    return "\n\n".join(parts), pub_ids[:5]


def _archive_worker(sid, user_id):
    job = ARCHIVE_JOBS.get(sid)
    db = SessionLocal()
    step = "初始化"
    current = ""
    try:
        sprint = db.query(models.Sprint).get(sid)
        project = db.query(models.Project).get(sprint.project_id)
        # D34：归档语料只取正式原型，demo 不进知识库
        publishes = [p for p in db.query(models.Publish).filter_by(sprint_id=sid).all()
                     if (p.kind or "prototype") == "prototype"]
        reqs = db.query(models.Requirement).filter_by(sprint_id=sid).all()

        # 步骤一：准备原型语料（该版本全部原型文字切块 + 向量）
        step = "准备原型语料"
        job["step"] = step
        proto_chunks = []  # [(publish_id, 片段文本)]，id 供 D21 条目关联原型
        for p in publishes:
            for c in chunk_text(p.extracted_text or "", 800):
                if c.strip():
                    proto_chunks.append((p.id, f"《{p.title}》：{c}"))
        proto_vecs = []
        if proto_chunks:
            try:
                proto_vecs = ai.embed([c for _, c in proto_chunks])
            except Exception as e:
                print(f"[warn] 归档原型语料向量化失败: {e}")
                proto_chunks = []

        def proto_retrieve(query, top=4):
            """按相关度返回 [(publish_id, 片段)]，供来源组装与条目关联原型记录。"""
            if not proto_chunks or not proto_vecs:
                return []
            try:
                qv = ai.embed([query])[0]
            except Exception:
                return []
            scored = sorted(
                zip(proto_chunks, proto_vecs),
                key=lambda cv: _cosine_sim(qv, cv[1]),
                reverse=True,
            )
            return [c for c, _ in scored[:top]]

        att_root = _attach_root()

        # 步骤二：逐需求生成知识条目（内容非空的已存在条目跳过，断点重试幂等；
        # 内容为空的条目视为上次失败产物，原地重生成——D25）
        step = "生成知识条目"
        job["step"] = step
        for r in reqs:
            current = r.title
            job["current"] = current
            existing = (
                db.query(models.KnowledgeEntry)
                .filter_by(sprint_id=sid, requirement_id=r.id, kind="requirement")
                .first()
            )
            if existing and (existing.content or "").strip():
                job["done"] += 1
                continue
            sources, pub_ids = _entry_sources(db, r, att_root, proto_retrieve)
            content = ai.chat(
                [
                    {"role": "system", "content": _ENTRY_PROMPT},
                    {
                        "role": "user",
                        "content": f"项目「{project.name if project else ''}」版本「{sprint.name}」的需求「{r.title}」资料如下：\n\n{sources}",
                    },
                ],
                max_tokens=16000,
            )
            if existing:
                entry = existing
                entry.content = content
                entry.edited = False
                entry.publish_ids = json.dumps(pub_ids)
                db.flush()
                knowledge.delete_entry(entry.id)  # 清掉空内容时代写入的旧向量
            else:
                entry = models.KnowledgeEntry(
                    project_id=project.id,
                    sprint_id=sid,
                    requirement_id=r.id,
                    req_no=r.req_no or "",
                    title=r.title,
                    kind="requirement",
                    content=content,
                    publish_ids=json.dumps(pub_ids),
                )
                db.add(entry)
                db.flush()
            chunks = _kb_chunk(
                f"需求「{r.title}」（{r.req_no or '-'}，版本 {sprint.name}）\n{content}"
            )
            if chunks:
                try:
                    embs = ai.embed(chunks)
                    knowledge.upsert_entry(
                        entry.id, project.id, sid, r.req_no or "", r.title, chunks, embs
                    )
                except Exception as e:
                    print(f"[warn] 知识条目向量化失败（entry {entry.id}）: {e}")
            db.commit()
            job["done"] += 1

        # 步骤三：生成版本索引（人工编辑过的索引沿用，不覆盖）
        step = "生成版本索引"
        current = ""
        job["step"] = step
        job["current"] = ""
        entries = (
            db.query(models.KnowledgeEntry)
            .filter_by(sprint_id=sid, kind="requirement")
            .order_by(models.KnowledgeEntry.id)
            .all()
        )
        idx_entry = (
            db.query(models.KnowledgeEntry).filter_by(sprint_id=sid, kind="index").first()
        )
        if idx_entry and idx_entry.edited:
            index_content = idx_entry.content
        else:
            listing = (
                "\n".join(
                    f"- {e.req_no or '-'} {e.title}：{(e.content or '')[:150]}"
                    for e in entries
                )
                or "（本版本无需求条目）"
            )
            template = (settings.archive.get("summary_template") or "").strip()
            instruction = (
                template.replace("{project}", project.name if project else "").replace(
                    "{version}", sprint.name
                )
                if template
                else _INDEX_PROMPT
            )
            index_content = ai.chat(
                [
                    {"role": "system", "content": instruction},
                    {
                        "role": "user",
                        "content": f"以下是项目「{project.name if project else ''}」版本「{sprint.name}」的需求知识条目清单：\n\n{listing}",
                    },
                ],
                max_tokens=16000,
            )
        if not (index_content or "").strip():
            # D25：空摘要不允许落库并置「已归档」，整任务失败可重试
            raise RuntimeError("版本索引生成结果为空")
        if not idx_entry or not idx_entry.edited:
            if not idx_entry:
                idx_entry = models.KnowledgeEntry(
                    project_id=project.id,
                    sprint_id=sid,
                    requirement_id=None,
                    req_no="",
                    title=f"{sprint.name} 版本索引",
                    kind="index",
                )
                db.add(idx_entry)
                db.flush()
            idx_entry.content = index_content
            idx_entry.publish_ids = json.dumps([p.id for p in publishes])
            db.commit()
            _kb_reembed(idx_entry, project.id, sprint.name)
        job["done"] += 1

        # 步骤四：联动状态
        step = "联动状态"
        job["step"] = step
        today = datetime.now().strftime("%Y-%m-%d")
        for r in reqs:
            if r.status != "已上线":
                r.status = "已上线"
            r.actual_online_date = r.actual_online_date or today  # 自愈重归档不覆盖原上线日期
        sprint.summary = index_content
        sprint.kb_doc_id = f"p{project.id}_s{sid}"
        sprint.status = "已归档"
        sprint.archived_by = user_id
        sprint.archived_at = _utcnow()
        db.commit()
        job["state"] = "done"
        job["step"] = "完成"
        job["done"] = job["total"]
    except Exception as e:
        db.rollback()
        job["state"] = "failed"
        job["error"] = (
            f"归档失败于步骤：{step}" + (f"（{current}）" if current else "") + f"：{e}"
        )
        print(f"[warn] 归档任务失败(sprint {sid}): {job['error']}")
    finally:
        db.close()


@app.post("/api/sprints/{sid}/archive")
def archive(sid: int, request: Request, user=Depends(get_user)):
    """启动归档后台任务（全体 PM，T2）：进度见 GET /archive-status；失败可按同接口重试。"""
    pm_or_admin(user)
    if not settings.ai_enabled:
        raise HTTPException(400, "请先在配置中心设置 LLM 网关")
    db = SessionLocal()
    sprint = db.query(models.Sprint).get(sid)
    if not sprint:
        db.close()
        raise HTTPException(404, "版本不存在")
    if sprint.status != "进行中":
        # D25 自愈通道：已归档但摘要为空（上次归档时模型返回空内容）允许重新归档
        if not (sprint.status == "已归档" and not (sprint.summary or "").strip()):
            db.close()
            raise HTTPException(400, "仅进行中的版本可归档")
    req_count = db.query(models.Requirement).filter_by(sprint_id=sid).count()
    db.close()
    existing = ARCHIVE_JOBS.get(sid)
    if existing and existing.get("state") == "running":
        raise HTTPException(400, "该版本归档任务正在进行中")
    total = req_count + 1  # 逐需求条目 + 版本索引
    ARCHIVE_JOBS[sid] = {
        "state": "running",
        "step": "初始化",
        "total": total,
        "done": 0,
        "current": "",
        "error": "",
    }
    threading.Thread(target=_archive_worker, args=(sid, user.id), daemon=True).start()
    return {"started": True, "total": total}


@app.get("/api/sprints/{sid}/archive-status")
def archive_status(sid: int, user=Depends(get_user)):
    job = ARCHIVE_JOBS.get(sid)
    if job:
        return dict(job)
    db = SessionLocal()
    sprint = db.query(models.Sprint).get(sid)
    db.close()
    if sprint and sprint.status == "已归档":
        return {"state": "done", "step": "完成", "total": 1, "done": 1, "current": "", "error": ""}
    return {"state": "idle"}


@app.get("/api/sprints/{sid}/report")
def get_report(sid: int, user=Depends(get_user)):
    """旧 reports 表只读兼容（历史备份，无新写入）。新版本摘要在 sprint.summary。"""
    db = SessionLocal()
    sprint = db.query(models.Sprint).get(sid)
    r = (
        db.query(models.Report)
        .filter_by(sprint_id=sid)
        .order_by(models.Report.id.desc())
        .first()
    )
    db.close()
    if sprint and sprint.summary:
        return {
            "exists": True,
            "content": sprint.summary,
            "created_at": str(sprint.archived_at or sprint.created_at),
        }
    if not r:
        return {"exists": False}
    return {
        "exists": True,
        "content": r.content,
        "created_at": str(r.created_at),
    }


# ---------- 知识库浏览与维护（D19；浏览=所有登录角色，编辑/重生成=PM/主账号） ----------
@app.get("/api/kb/entries")
def kb_entries(
    project_id: int = 0,
    sprint_id: int = 0,
    q: str = "",
    user=Depends(get_user),
):
    db = SessionLocal()
    query = db.query(models.KnowledgeEntry)
    if project_id:
        query = query.filter(models.KnowledgeEntry.project_id == project_id)
    if sprint_id:
        query = query.filter(models.KnowledgeEntry.sprint_id == sprint_id)
    entries = query.order_by(
        models.KnowledgeEntry.sprint_id.desc(), models.KnowledgeEntry.id
    ).all()
    sprint_names = {
        s.id: s.name
        for s in db.query(models.Sprint)
        .filter(models.Sprint.id.in_({e.sprint_id for e in entries} or {0}))
        .all()
    }
    out = []
    needle = q.strip().lower()
    for e in entries:
        if needle and not (
            needle in (e.title or "").lower()
            or needle in (e.req_no or "").lower()
            or needle in (e.content or "").lower()
        ):
            continue
        out.append(
            {
                "id": e.id,
                "project_id": e.project_id,
                "sprint_id": e.sprint_id,
                "sprint_name": sprint_names.get(e.sprint_id, ""),
                "requirement_id": e.requirement_id,
                "req_no": e.req_no or "",
                "title": e.title or "",
                "kind": e.kind,
                "edited": bool(e.edited),
                "preview": (e.content or "")[:120],
                "updated_at": str(e.updated_at or e.created_at or ""),
            }
        )
    db.close()
    return {"entries": out}


@app.get("/api/kb/entries/{eid}")
def kb_entry_detail(eid: int, user=Depends(get_user)):
    db = SessionLocal()
    e = db.query(models.KnowledgeEntry).get(eid)
    if not e:
        db.close()
        raise HTTPException(404, "知识条目不存在")
    sprint = db.query(models.Sprint).get(e.sprint_id)
    data = {
        "id": e.id,
        "project_id": e.project_id,
        "sprint_id": e.sprint_id,
        "sprint_name": sprint.name if sprint else "",
        "requirement_id": e.requirement_id,
        "req_no": e.req_no or "",
        "title": e.title or "",
        "kind": e.kind,
        "edited": bool(e.edited),
        "content": e.content or "",
        "updated_at": str(e.updated_at or e.created_at or ""),
    }
    db.close()
    return data


def _kb_reembed(entry, project_id, sprint_name):
    """条目内容变更后重建向量切片（先删后写，失败不阻断）。"""
    knowledge.delete_entry(entry.id)
    label = (
        f"版本索引「{sprint_name}」"
        if entry.kind == "index"
        else f"需求「{entry.title}」（{entry.req_no or '-'}，版本 {sprint_name}）"
    )
    chunks = _kb_chunk(f"{label}\n{entry.content}")
    if not chunks:
        return
    try:
        embs = ai.embed(chunks)
        knowledge.upsert_entry(
            entry.id,
            project_id,
            entry.sprint_id,
            entry.req_no or "",
            entry.title or "",
            chunks,
            embs,
        )
    except Exception as e:
        print(f"[warn] 知识条目重建向量失败（entry {entry.id}）: {e}")


@app.put("/api/kb/entries/{eid}")
def kb_entry_update(eid: int, content: str = Form(...), user=Depends(get_user)):
    pm_or_admin(user)
    if not content.strip():
        raise HTTPException(400, "内容不能为空")
    db = SessionLocal()
    e = db.query(models.KnowledgeEntry).get(eid)
    if not e:
        db.close()
        raise HTTPException(404, "知识条目不存在")
    e.content = content.strip()
    e.edited = True
    db.commit()
    sprint = db.query(models.Sprint).get(e.sprint_id)
    _kb_reembed(e, e.project_id, sprint.name if sprint else "")
    db.close()
    return {"ok": True}


@app.post("/api/kb/entries/{eid}/regenerate")
def kb_entry_regenerate(eid: int, user=Depends(get_user)):
    """用当前来源（字段+材料+对话+原型片段）重新生成条目内容，覆盖人工修改。"""
    pm_or_admin(user)
    if not settings.ai_enabled:
        raise HTTPException(400, "请先在配置中心设置 LLM 网关")
    db = SessionLocal()
    e = db.query(models.KnowledgeEntry).get(eid)
    if not e:
        db.close()
        raise HTTPException(404, "知识条目不存在")
    if e.kind != "requirement":
        db.close()
        raise HTTPException(400, "版本索引请通过重新归档生成")
    req = db.query(models.Requirement).get(e.requirement_id)
    sprint = db.query(models.Sprint).get(e.sprint_id)
    project = db.query(models.Project).get(e.project_id)
    if not req or not sprint:
        db.close()
        raise HTTPException(400, "关联需求或版本已不存在，无法重新生成")
    publishes = db.query(models.Publish).filter_by(sprint_id=e.sprint_id).all()
    proto_chunks = []  # [(publish_id, 片段文本)]，id 供 D21 条目关联原型
    for p in publishes:
        for c in chunk_text(p.extracted_text or "", 800):
            if c.strip():
                proto_chunks.append((p.id, f"《{p.title}》：{c}"))
    proto_vecs = []
    if proto_chunks:
        try:
            proto_vecs = ai.embed([c for _, c in proto_chunks])
        except Exception as ex:
            print(f"[warn] 条目重生成原型语料向量化失败: {ex}")
            proto_chunks = []

    def proto_retrieve(query, top=4):
        if not proto_chunks or not proto_vecs:
            return []
        try:
            qv = ai.embed([query])[0]
        except Exception:
            return []
        scored = sorted(
            zip(proto_chunks, proto_vecs),
            key=lambda cv: _cosine_sim(qv, cv[1]),
            reverse=True,
        )
        return [c for c, _ in scored[:top]]

    try:
        sources, pub_ids = _entry_sources(db, req, _attach_root(), proto_retrieve)
        content = ai.chat(
            [
                {"role": "system", "content": _ENTRY_PROMPT},
                {
                    "role": "user",
                    "content": f"项目「{project.name if project else ''}」版本「{sprint.name}」的需求「{req.title}」资料如下：\n\n{sources}",
                },
            ],
            max_tokens=16000,
        )
    except Exception as ex:
        db.close()
        print(f"[warn] 知识条目重生成失败: {ex}")
        raise HTTPException(502, f"重新生成失败：{ex}")
    e.content = content
    e.edited = False
    e.publish_ids = json.dumps(pub_ids)
    db.commit()
    _kb_reembed(e, e.project_id, sprint.name)
    data = {"ok": True, "content": e.content}
    db.close()
    return data


# ---------- 知识库查询问答（内置问答 UI，D1；所有登录角色可用） ----------
@app.post("/api/query")
def query(project_id: int = Form(...), question: str = Form(...), user=Depends(get_user)):
    if not settings.ai_enabled:
        raise HTTPException(400, "请先在配置中心设置 LLM 网关")
    if not question.strip():
        raise HTTPException(400, "请输入问题")

    qvec = ai.embed([question])[0]
    hits = knowledge.query_knowledge(project_id, qvec, with_meta=True)
    docs = [
        f"《{(m or {}).get('title') or '知识条目'}》（{(m or {}).get('req_no') or '版本索引'}）：\n{d}"
        for d, m in hits
    ]
    context = "\n\n".join(docs) if docs else "（该项目的知识库中暂无相关记录）"
    try:
        answer = ai.chat(
            [
                {
                    "role": "system",
                    "content": "你是团队知识库助手。请仅基于下方「知识库内容」回答用户问题；若资料不足，请明确说明。回答用中文，并尽量引用相关迭代信息。",
                },
                {
                    "role": "user",
                    "content": f"知识库内容：\n{context}\n\n用户问题：{question}",
                },
            ]
        )
    except Exception as e:
        print(f"[warn] 知识库问答模型调用失败: {e}")
        raise HTTPException(502, "模型调用失败，请稍后重试或联系主账号")

    # D21：按命中条目关联的原型（publish_ids）给出佐证直链，按命中顺序去重；已删原型自动剔除
    prototypes = []
    entry_ids = [int((m or {}).get("entry_id") or 0) for _, m in hits]
    entry_ids = [i for i in dict.fromkeys(entry_ids) if i]
    if entry_ids:
        db = SessionLocal()
        try:
            entries = (
                db.query(models.KnowledgeEntry)
                .filter(models.KnowledgeEntry.id.in_(entry_ids))
                .all()
            )
            by_id = {e.id: e for e in entries}
            pub_ids = []
            for eid in entry_ids:  # 保持检索命中顺序
                e = by_id.get(eid)
                if not e or not e.publish_ids:
                    continue
                try:
                    ids = json.loads(e.publish_ids)
                except (TypeError, ValueError):
                    continue
                for pid in ids:
                    if isinstance(pid, int) and pid not in pub_ids:
                        pub_ids.append(pid)
            if pub_ids:
                pubs = (
                    db.query(models.Publish)
                    .filter(models.Publish.id.in_(pub_ids))
                    .all()
                )
                pmap = {p.id: p for p in pubs}
                snames = {
                    s.id: s.name
                    for s in db.query(models.Sprint)
                    .filter(models.Sprint.id.in_({p.sprint_id for p in pubs} or {0}))
                    .all()
                }
                for pid in pub_ids:
                    p = pmap.get(pid)
                    if not p:
                        continue  # 原型已删除
                    prototypes.append(
                        {
                            "title": p.title or "未命名原型",
                            "sprint_name": snames.get(p.sprint_id, ""),
                            "url": f"/p/{p.access_key}/{p.entry_path or ''}",
                        }
                    )
        finally:
            db.close()
    return JSONResponse({"answer": answer, "sources": docs, "prototypes": prototypes})


# ---------- 统计分析 ----------
@app.get("/api/stats")
def stats(
    from_: str = Query("", alias="from"),
    to: str = "",
    pm_id: int = 0,
    project_id: int = 0,
    sprint_id: int = 0,
    user=Depends(get_user),
):
    """R1 三维度计数（项目/状态/PM）+ R2 PM 工作量表。
    筛选（D11）：requirement.created_at 范围（默认全部），可叠加 PM / 项目 / 迭代
    （sprint_id=-1 表示"未纳入迭代"）。"""
    if from_ and to and from_ > to:
        raise HTTPException(400, "时间范围不正确")
    db = SessionLocal()
    q = db.query(models.Requirement)
    if from_:
        q = q.filter(func.date(models.Requirement.created_at) >= from_)
    if to:
        q = q.filter(func.date(models.Requirement.created_at) <= to)
    if pm_id:
        q = q.filter(models.Requirement.pm_id == pm_id)
    if project_id:
        q = q.filter(models.Requirement.project_id == project_id)
    if sprint_id == -1:
        q = q.filter(models.Requirement.sprint_id.is_(None))
    elif sprint_id:
        q = q.filter(models.Requirement.sprint_id == sprint_id)
    reqs = q.all()

    by_project = {}
    by_status = {}
    by_pm = {}
    for r in reqs:
        proj = db.query(models.Project).get(r.project_id)
        pname = proj.name if proj else f"项目{r.project_id}"
        by_project[pname] = by_project.get(pname, 0) + 1
        by_status[r.status] = by_status.get(r.status, 0) + 1
        pm = db.query(models.User).get(r.pm_id) if r.pm_id else None
        pmname = _user_name(pm) if pm else "未指派"
        by_pm[pmname] = by_pm.get(pmname, 0) + 1

    # R2：PM 工作量（需求计数，D26：人天字段已移除）
    workload = {}
    for r in reqs:
        pm = db.query(models.User).get(r.pm_id) if r.pm_id else None
        pmname = _user_name(pm) if pm else "未指派"
        w = workload.setdefault(pmname, {"count": 0})
        w["count"] += 1
    r2_rows = [
        {"pm": k, "count": v["count"]}
        for k, v in sorted(workload.items())
    ]
    r2_total = {
        "pm": "合计",
        "count": sum(x["count"] for x in r2_rows),
    }
    db.close()
    return {
        "r1": {
            "by_project": [{"name": k, "count": v} for k, v in sorted(by_project.items())],
            "by_status": [{"name": k, "count": v} for k, v in by_status.items()],
            "by_pm": [{"name": k, "count": v} for k, v in sorted(by_pm.items())],
        },
        "r2": {"rows": r2_rows, "total": r2_total},
        "total": len(reqs),
    }


# ---------- 系统配置（仅主账号，D4：config.yaml + Settings） ----------
@app.get("/config")
def config_page(request: Request, user=Depends(admin_user)):
    # 配置中心已并入 SPA（#/pc/admin），旧路由跳转
    return RedirectResponse("/#/pc/admin", status_code=303)


@app.get("/api/config")
def get_config(user=Depends(admin_user)):
    return {
        "host": settings.host,
        "port": settings.port,
        "secret_key": settings.secret_key,
        "publish_root": settings.publish_root,
        "ai_base_url": settings.ai_base_url,
        "ai_api_key": settings.ai_api_key,
        "chat_model": settings.chat_model,
        "embedding_model": settings.embedding_model,
        "ai_enabled": settings.ai_enabled,
        "embedding_provider": settings.embedding_provider,
        "embedding_model_path": settings.embedding_model_path,
        "design_enabled": settings.design_enabled,
        "design_agent": settings.design_agent,
        "agent2": settings.agent2,
        "archive": settings.archive,
        "knowledge": settings.knowledge,
        "chroma_dir": settings.chroma_dir,
        "dingtalk": {
            **settings.dingtalk,
            "app_secret": "***" if settings.dingtalk.get("app_secret") else "",
        },
        "escalate": settings.escalate,
        "jira": {
            **settings.jira,
            "api_token": "***" if settings.jira.get("api_token") else "",
        },
    }


@app.post("/api/config")
def post_config(
    section: str = Form("all"),  # basic/models/design_agent/agent2/archive/all
    host: str = Form(""),
    port: str = Form(""),
    secret_key: str = Form(""),
    publish_root: str = Form(""),
    ai_base_url: str = Form(""),
    ai_api_key: str = Form(""),
    chat_model: str = Form(""),
    embedding_model_path: str = Form(""),
    design_description: str = Form(""),
    design_base_url: str = Form(""),
    design_api_key: str = Form(""),
    design_model: str = Form(""),
    design_use_kb: str = Form(""),
    design_max_iterations: str = Form(""),
    design_max_input_length: str = Form(""),
    design_temperature: str = Form(""),
    design_top_p: str = Form(""),
    design_max_tokens: str = Form(""),
    design_skills_json: str = Form(""),
    agent2_system_prompt: str = Form(""),
    agent2_use_kb: str = Form(""),
    archive_summary_template: str = Form(""),
    kb_top_k: str = Form(""),
    kb_chunk_size: str = Form(""),
    kb_chunk_overlap: str = Form(""),
    dt_app_key: str = Form(""),
    dt_app_secret: str = Form(""),
    dt_agent_id: str = Form(""),
    dt_hub_base_url: str = Form(""),
    escalate_ac_template: str = Form(""),
    jr_base_url: str = Form(""),
    jr_email: str = Form(""),
    jr_api_token: str = Form(""),
    jr_project_key: str = Form(""),
    jr_story_template: str = Form(""),
    user=Depends(admin_user),
):
    updates = {}
    if section in ("basic", "all"):
        updates.update({
            "host": host or "",
            "port": int(port) if port else "",
            "secret_key": secret_key or "",
            "publish_root": publish_root or "",
        })
    if section in ("models", "all"):
        updates["ai"] = {
            "base_url": ai_base_url or "",
            "api_key": ai_api_key or "",
            "chat_model": chat_model or "",
        }
        if embedding_model_path.strip():
            updates["embedding"] = {
                "provider": "local_onnx",
                "model_path": embedding_model_path.strip(),
            }
    if section in ("design_agent", "all"):
        da = {}
        da["description"] = design_description or ""
        da["base_url"] = design_base_url or ""
        da["api_key"] = design_api_key or ""
        da["chat_model"] = design_model or ""
        if design_use_kb != "":
            da["use_kb"] = design_use_kb.strip().lower() in ("1", "true", "on")
        for field, raw, cast, default in (
            ("max_iterations", design_max_iterations, int, 1),
            ("max_input_length", design_max_input_length, int, 24000),
            ("temperature", design_temperature, float, 0.7),
            ("top_p", design_top_p, float, 0.9),
            ("max_tokens", design_max_tokens, int, 3000),
        ):
            try:
                da[field] = cast(raw) if raw != "" else default
            except (TypeError, ValueError):
                raise HTTPException(400, f"高级选项 {field} 格式不正确")
        if da["max_iterations"] < 1:
            da["max_iterations"] = 1
        if design_skills_json.strip():
            try:
                skills = json.loads(design_skills_json)
            except json.JSONDecodeError:
                raise HTTPException(400, "技能配置格式不正确")
            cleaned = []
            for s in skills:
                key = str(s.get("key", "")).strip()
                name = str(s.get("name", "")).strip()
                prompt = str(s.get("prompt", "")).strip()
                if not key or not name or not prompt:
                    continue
                cleaned.append({
                    "key": key, "name": name, "prompt": prompt,
                    "enabled": bool(s.get("enabled", True)),
                })
            da["skills"] = cleaned
        updates["design_agent"] = da
    if section in ("agent2", "all"):
        a2 = {}
        a2["system_prompt"] = agent2_system_prompt or ""
        if agent2_use_kb != "":
            a2["use_kb"] = agent2_use_kb.strip().lower() in ("1", "true", "on")
        updates["agent2"] = a2
    if section in ("archive", "all"):
        updates["archive"] = {"summary_template": archive_summary_template or ""}
    if section in ("knowledge", "all"):
        try:
            top_k = int(kb_top_k) if kb_top_k.strip() else 5
            chunk_size = int(kb_chunk_size) if kb_chunk_size.strip() else 1200
            chunk_overlap = int(kb_chunk_overlap) if kb_chunk_overlap.strip() else 0
        except ValueError:
            raise HTTPException(400, "知识库参数必须是整数")
        if not (1 <= top_k <= 20):
            raise HTTPException(400, "top_k 取值范围 1-20")
        if not (200 <= chunk_size <= 4000):
            raise HTTPException(400, "chunk_size 取值范围 200-4000")
        if not (0 <= chunk_overlap <= chunk_size // 2):
            raise HTTPException(400, "chunk_overlap 需在 0 到 chunk_size 的一半之间")
        updates["knowledge"] = {
            "top_k": top_k,
            "chunk_size": chunk_size,
            "chunk_overlap": chunk_overlap,
        }
    if section in ("dingtalk", "all"):
        updates["dingtalk"] = {
            "app_key": dt_app_key.strip(),
            # app_secret 不回显：留空表示沿用已保存值
            "app_secret": dt_app_secret.strip() or settings.dingtalk.get("app_secret") or "",
            "agent_id": dt_agent_id.strip(),
            "hub_base_url": dt_hub_base_url.strip(),
        }
    if section in ("escalate", "all"):
        updates["escalate"] = {"ac_template": escalate_ac_template or ""}
    if section == "jira_tpl":
        updates["jira"] = {"story_template": jr_story_template or ""}
    if section in ("jira", "all"):
        cur = dict(settings.jira or {})
        cur["base_url"] = jr_base_url.strip()
        cur["email"] = jr_email.strip()
        # api_token 不回显：留空表示沿用已保存值
        cur["api_token"] = jr_api_token.strip() or cur.get("api_token") or ""
        cur["project_key"] = jr_project_key.strip()
        if jr_story_template.strip():
            cur["story_template"] = jr_story_template
        updates["jira"] = cur
    save_config(updates)
    return {"ok": True}


@app.get("/api/config/disk-usage")
def disk_usage(user=Depends(admin_user)):
    """R3 磁盘用量（D12 修订）：已用量=SUM(publishes.total_size)+SUM(attachments.size)；
    总容量优先 DISK_TOTAL_BYTES 环境变量（手动覆盖），未设置时自动探测 publish_root
    所在磁盘的真实总容量；附带该磁盘真实剩余空间。"""
    db = SessionLocal()
    proto_bytes = db.query(func.coalesce(func.sum(models.Publish.total_size), 0)).scalar()
    att_bytes = db.query(func.coalesce(func.sum(models.RequirementAttachment.size), 0)).scalar()
    db.close()
    root = settings.publish_root
    os.makedirs(root, exist_ok=True)
    du = shutil.disk_usage(root)
    env_total = os.environ.get("DISK_TOTAL_BYTES")
    auto_total = not env_total
    total = int(env_total) if env_total else int(du.total)
    used = int(proto_bytes) + int(att_bytes)
    ratio = used / total if total else 0
    return {
        "proto_bytes": int(proto_bytes),
        "attachment_bytes": int(att_bytes),
        "used_bytes": used,
        "total_bytes": total,
        "used_gb": round(used / 1024 ** 3, 1),
        "total_gb": round(total / 1024 ** 3, 1),
        "free_gb": round(du.free / 1024 ** 3, 1),
        "auto_total": auto_total,
        "percent": round(ratio * 100, 2),
        "over_80": ratio >= 0.8,
    }


@app.post("/api/config/test-llm")
def test_llm(
    base_url: str = Form(""),
    api_key: str = Form(""),
    model: str = Form(""),
    user=Depends(admin_user),
):
    """LLM 连通性测试（D13）：表单值优先、留空回落已保存配置；返回分步日志，
    失败时按异常类型给出排查建议。"""
    t0 = time.monotonic()
    steps = []

    def done(ok, error="", hint=""):
        return {
            "ok": ok, "steps": steps, "duration_ms": int((time.monotonic() - t0) * 1000),
            "error": error, "hint": hint,
        }

    bu = (base_url.strip() or settings.ai_base_url or "").rstrip("/")
    ak = api_key.strip() or settings.ai_api_key or ""
    md = model.strip() or settings.chat_model or ""

    missing = [n for n, v in [("base_url", bu), ("api_key", ak), ("model", md)] if not v]
    if missing:
        steps.append({"name": "配置检查", "ok": False, "detail": "缺少：" + " / ".join(missing)})
        return done(False, "配置不完整", "请先填写并保存 LLM 网关三项配置（base_url / api_key / model）")
    masked = (ak[:6] + "…" + ak[-4:]) if len(ak) > 10 else "***"
    steps.append({"name": "配置检查", "ok": True,
                  "detail": f"base_url={bu} · model={md} · api_key={masked}"})

    url = f"{bu}/chat/completions"
    steps.append({"name": "发送测试请求", "ok": True, "detail": f"POST {url}（超时上限 30s）"})
    try:
        r = requests.post(
            url,
            json={"model": md, "messages": [{"role": "user", "content": "连通性测试，请只回复两个字：正常"}], "max_tokens": 16},
            headers={"Authorization": f"Bearer {ak}", "Content-Type": "application/json"},
            timeout=30,
        )
        ms = int((time.monotonic() - t0) * 1000)
    except requests.exceptions.ConnectTimeout:
        steps.append({"name": "网关响应", "ok": False, "detail": "连接超时（30s 内未建立连接）"})
        return done(False, "连接超时", "检查 base_url 是否写对、目标服务是否在线、本机到网关的网络/代理是否通")
    except requests.exceptions.ReadTimeout:
        steps.append({"name": "网关响应", "ok": False, "detail": "响应超时（已连接，30s 内未返回数据）"})
        return done(False, "响应超时", "网关能连上但响应过慢，检查模型服务负载或换一个模型试试")
    except requests.exceptions.SSLError:
        steps.append({"name": "网关响应", "ok": False, "detail": "HTTPS 证书校验失败"})
        return done(False, "证书校验失败", "检查网关证书是否有效；内网自签证书可改用 http:// 或配置受信证书")
    except requests.exceptions.ConnectionError as e:
        steps.append({"name": "网关响应", "ok": False, "detail": f"无法建立连接：{str(e)[:120]}"})
        return done(False, "连接失败", "域名解析失败或端口不通，检查 base_url 的域名/端口是否正确")
    except Exception as e:
        steps.append({"name": "网关响应", "ok": False, "detail": f"请求异常：{str(e)[:120]}"})
        return done(False, "请求异常", "把上面的异常信息发给网关运维排查")

    if r.status_code != 200:
        excerpt = (r.text or "").strip().replace("\n", " ")[:200]
        steps.append({"name": "网关响应", "ok": False, "detail": f"HTTP {r.status_code} · {excerpt}"})
        hint = {
            400: "请求格式被拒，确认模型名是否正确、网关是否兼容 OpenAI chat/completions 格式",
            401: "api_key 无效或已过期，重新生成一个试试",
            403: "api_key 没有该模型的访问权限",
            404: "路径不存在，确认 base_url 以 /v1 结尾（不要带 /chat/completions）",
            429: "触发限流或账户额度不足",
        }.get(r.status_code, "网关返回错误，把 HTTP 状态码和响应内容发给网关运维排查")
        return done(False, f"HTTP {r.status_code}", hint)

    steps.append({"name": "网关响应", "ok": True, "detail": f"HTTP 200 · 耗时 {ms} ms"})
    try:
        content = r.json()["choices"][0]["message"]["content"]
    except Exception:
        steps.append({"name": "解析模型回复", "ok": False,
                      "detail": f"HTTP 200 但响应不是标准 OpenAI 格式：{(r.text or '')[:160]}"})
        return done(False, "响应格式不兼容", "网关需兼容 OpenAI chat/completions 响应格式（choices[0].message.content）")
    excerpt = (content or "").strip()[:80]
    if not excerpt:
        # 推理型模型可能只回 reasoning_content，content 为空也算协议正常
        alt = (r.json()["choices"][0]["message"].get("reasoning_content") or "").strip()
        excerpt = ("（content 为空，模型推理摘录：" + alt[:60] + "…）") if alt else "（content 为空，但响应协议正常，判定为连通）"
    steps.append({"name": "解析模型回复", "ok": True, "detail": f"模型回复：{excerpt}"})
    return done(True)


@app.post("/api/config/test-dingtalk")
def test_dingtalk(userid: str = Form(""), user=Depends(admin_user)):
    """钉钉推送测试（D22）：分步日志（检查配置→检查接收人→获取 token→发送通知）。"""
    return dingtalk.test_push(userid)


@app.post("/api/config/test-jira")
def test_jira(user=Depends(admin_user)):
    """Jira 连接测试（D23）：分步日志（检查配置→认证→项目校验），不真建 issue。"""
    return jira.test_connection()
